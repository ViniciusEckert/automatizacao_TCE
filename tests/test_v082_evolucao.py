import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from pypdf import PdfReader
from app import create_app
from src.persistencia.arquivos import ArquivoBrutoStore

class Evolucao082Test(unittest.TestCase):
    def setUp(self):
        self.c=create_app().test_client()
    def test_painel_modelo_e_acumulada(self):
        hist=self.c.get('/api/amostras/curitiba-historico').json
        r=self.c.post('/api/exportar',json=hist);self.assertEqual(r.status_code,200)
        w=load_workbook(BytesIO(r.data));v=load_workbook(BytesIO(r.data),data_only=True);p=w['Painel financeiro']
        self.assertEqual(p['A1'].value,'PAINEL ORÇAMENTÁRIO/FISCAL')
        self.assertEqual(p['B4'].value,2019)
        self.assertIn('$B$6',p['C7'].value)
        self.assertIn('C6/B6',p['C8'].value)
        self.assertEqual(p['B4'].fill.fgColor.rgb[-6:],'000000')
        self.assertEqual(len(p._charts),1);self.assertEqual(len(p._charts[0].series),3)
        self.assertAlmostEqual(v['Painel financeiro']['C7'].value,hist['resultados'][1]['receita_corrente_liquida_ajustada']/hist['resultados'][0]['receita_corrente_liquida_ajustada']-1)
        self.assertIn('RGF',w['Parâmetros do painel']['B5'].value)
    def test_pdf_profissao_generica_e_todos_os_valores(self):
        dados=self.c.get('/api/salarios/amostras/curitiba-administrativo-2026').json
        r=self.c.post('/api/salarios/relatorio-pdf',json=dados);self.assertEqual(r.status_code,200,r.json)
        text=' '.join(' '.join(p.extract_text().split()) for p in PdfReader(BytesIO(r.data)).pages)
        self.assertIn('Auxiliar Administrativo',text);self.assertNotIn('PSPN nacional',text)
        self.assertIn('11.490,21',text)
    def test_impressao_todas_as_abas_salariais(self):
        d=self.c.get('/api/salarios/amostras/cafezal-do-sul-2026').json
        r=self.c.post('/api/salarios/exportar-excel',json=d)
        for ws in load_workbook(BytesIO(r.data)):
            self.assertEqual(ws.page_setup.orientation,'landscape');self.assertEqual(ws.page_setup.fitToWidth,1);self.assertTrue(ws.print_area)
    def test_coletas_versionadas_sem_perder_caminho_legado(self):
        with TemporaryDirectory() as t:
            store=ArquivoBrutoStore(t)
            a=store.salvar('1','2',2025,'20','primeira')
            b=store.salvar('1','2',2025,'20','segunda')
            self.assertEqual(a,b);self.assertEqual(b.read_text(),'segunda')
            versions=list(Path(t).rglob('versoes/*.csv'));self.assertEqual(len(versions),2)
            self.assertEqual({p.read_text() for p in versions},{'primeira','segunda'})
