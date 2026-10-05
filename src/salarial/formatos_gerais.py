"""Conversões locais e OCR opcional, sem buscar URLs ou executar macros."""
from pathlib import Path
from io import BytesIO
import shutil,subprocess,tempfile
import pdfplumber
from docx import Document
from src.salarial.leitor_geral import ler_matrizes,ler_texto
from src.salarial.documento_service import _texto_doc_legado


def ocr_pagina(page):
    exe=shutil.which('tesseract')
    if not exe:raise ValueError('PDF imagem: instale Tesseract OCR ou envie PDF com texto/Excel.')
    with tempfile.TemporaryDirectory(prefix='jornada-ocr-') as tmp:
        p=Path(tmp)/'pagina.png';page.to_image(resolution=180).original.save(p)
        langs=subprocess.run([exe,'--list-langs'],capture_output=True,text=True,timeout=10).stdout.splitlines()
        lang='por' if 'por' in langs else 'eng'
        r=subprocess.run([exe,str(p),'stdout','-l',lang,'--psm','6'],capture_output=True,text=True,timeout=45)
        if r.returncode:raise ValueError('OCR não conseguiu ler a página; transcreva na revisão.')
        return r.stdout


def ler_pdf(nome,conteudo,organizar,ignorar_paginas=()):
    out=[];pend=[];textos=[]
    with pdfplumber.open(BytesIO(conteudo)) as doc:
        if len(doc.pages)>300:raise ValueError('Envie recorte de até 300 páginas.')
        ocr_count=0
        for i,page in enumerate(doc.pages,1):
            if i in ignorar_paginas:continue
            text=page.extract_text() or '';ocr=False
            if len(text.strip())<20:
                ocr_count+=1
                if ocr_count>10:pend.append(f'Página {i}: limite de 10 páginas OCR por arquivo; envie um recorte.');continue
                try:text=ocr_pagina(page);ocr=True
                except (ValueError,subprocess.TimeoutExpired) as e:pend.append(f'Página {i}: {e}');continue
            textos.append(f'PÁGINA {i}\n{text}')
            d,p=ler_texto(text,f'{nome}, página {i}',organizar)
            if not d and not p:
                d,p=ler_matrizes(page.extract_tables(),f'{nome}, página {i}',organizar,text)
            if ocr:
                pend.append(f'Página {i}: OCR aplicado; confira cada valor no original. Não há garantia de acurácia.')
                for c in d:c['avisos_importacao'].append('Valores lidos por OCR: revisão de todas as células obrigatória.');c['origem_ocr']=True
            out.extend(d);pend.extend(p)
            if not d:pend.append(f'Página {i}: não foi reconhecida uma tabela completa; consulte a revisão manual.')
    return out,'\n'.join(textos),pend


def ler_word(nome,conteudo,organizar):
    ext=Path(nome).suffix.lower()
    if ext=='.doc':
        exe=shutil.which('soffice') or shutil.which('libreoffice')
        if exe:
            with tempfile.TemporaryDirectory(prefix='jornada-word-') as tmp:
                pasta=Path(tmp);src=pasta/'original.doc';src.write_bytes(conteudo)
                proc=subprocess.run([exe,f'-env:UserInstallation={(pasta / "perfil").as_uri()}', '--headless','--convert-to','docx','--outdir',str(pasta),str(src)],capture_output=True,timeout=45)
                convertido=pasta/'original.docx'
                if proc.returncode==0 and convertido.exists():
                    d,t,p=ler_word('convertido.docx',convertido.read_bytes(),organizar)
                    for c in d:c['arquivo_fonte']=c['arquivo_fonte'].replace('convertido.docx',nome);c['avisos_importacao'].append('DOC convertido localmente para DOCX; confira a conversão no original.')
                    return d,t,p
        text=_texto_doc_legado(conteudo)
        d,p=ler_texto(text,nome,organizar);return d,text,p
    try:doc=Document(BytesIO(conteudo))
    except Exception as erro:raise ValueError('DOCX inválido ou protegido; envie um documento válido.') from erro
    text='\n'.join(p.text for p in doc.paragraphs)
    d,p=ler_matrizes([[[c.text for c in r.cells] for r in t.rows] for t in doc.tables],nome,organizar,text)
    if not d and not p:d,p=ler_texto(text,nome,organizar)
    return d,text,p


def ler_xls(conteudo):
    import xlrd
    try:book=xlrd.open_workbook(file_contents=conteudo)
    except xlrd.XLRDError as erro:raise ValueError('Arquivo XLS inválido ou protegido; envie XLSX/CSV.') from erro
    tables=[]
    for sheet in book.sheets():
        if sheet.nrows>20000 or sheet.ncols>220:raise ValueError('Planilha extensa: envie recorte.')
        tables.append([sheet.row_values(i) for i in range(sheet.nrows)])
    return tables


def ler_imagem(nome,conteudo,organizar):
    from PIL import Image
    from types import SimpleNamespace
    try:im=Image.open(BytesIO(conteudo))
    except Exception as erro:raise ValueError('Imagem inválida; envie PNG/JPG legível.') from erro
    if im.width*im.height>20000000:raise ValueError('Imagem extensa: envie um recorte de até 20 megapixels.')
    class Pagina:
        def to_image(self,resolution):return SimpleNamespace(original=im.convert('RGB'))
    texto=ocr_pagina(Pagina())
    d,p=ler_texto(texto,nome,organizar)
    for c in d:c['origem_ocr']=True;c['avisos_importacao'].append('Imagem lida por OCR: confira cada célula no original.')
    p.append('OCR de imagem: confirmação dos valores obrigatória; leitura pode confundir letras e dígitos.')
    return d,texto,p
