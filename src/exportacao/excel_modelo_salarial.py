"""Adaptador do exportador Python existente para o modelo do professor."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from src.exportacao.modelo_salarial import criar_modelo
from src.salarial.analise_service import analisar_tabela_salarial


def normalizar_analise(dados):
    if not isinstance(dados, dict):
        raise ValueError("Envie os dados da análise.")
    if "entrada" in dados:
        return analisar_tabela_salarial(dados["entrada"])
    if "niveis" in dados:
        return analisar_tabela_salarial(dados)
    raise ValueError("Recalcule a análise antes de exportar.")


def gerar_excel_modelo(dados):
    spec = criar_modelo(normalizar_analise(dados))
    wb = Workbook(); wb.remove(wb.active)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    for s in spec['sheets']:
        ws = wb.create_sheet(s['name']); ws.sheet_view.showGridLines = False
        for j,width in enumerate(s['widths'],1):ws.column_dimensions[get_column_letter(j)].width=width
        for r,h in s['heights'].items():ws.row_dimensions[int(r)].height=h
        for item in s['cells']:
            c = ws[item['address']]
            c.value = item.get('formula',item['value'])
            if not item.get('formula') and isinstance(c.value,str):c.data_type='s'
            c.font=Font(name='Arial',size=10)
            c.alignment=Alignment(vertical='center',horizontal='right' if isinstance(c.value,(int,float)) or item.get('formula') else 'left')
            if 'numberFormat' in item:c.number_format=item['numberFormat']
        for r in s['merges']:ws.merge_cells(r)
        for f in s['formats']:
            for line in ws[f['range']]:
                for c in line:
                    c.font=Font(name='Arial',size=f.get('size',10),bold=f.get('bold',False),color=f.get('fontColor','#000000').lstrip('#'))
                    c.alignment=Alignment(horizontal=f.get('align',c.alignment.horizontal),vertical='center',wrap_text=f.get('wrap',False))
                    if f.get('fill'):c.fill=PatternFill('solid',fgColor=f['fill'].lstrip('#'))
                    if f.get('bottom'):c.border=Border(bottom=Side(style='thin',color=f['bottom'].lstrip('#')))
        for rule in s['conditional']:
            ws.conditional_formatting.add(rule['range'],FormulaRule(formula=[rule['formula'].lstrip('=')],fill=PatternFill('solid',fgColor=rule['fill'].lstrip('#'))))
        ws.freeze_panes='D4' if s['name'] in {'Comparação salarial','Dados originais'} else 'B3'
        if not s.get('print'):
            ws.print_area=ws.calculate_dimension();ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
            ws.print_options.horizontalCentered=True
        if s.get('print'):
            ws.print_area=s['print'];ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
            ws.print_options.horizontalCentered=True
    dv=DataValidation(type='list',formula1='"arredondar,truncar"');wb['Parâmetros'].add_data_validation(dv);dv.add(wb['Parâmetros']['B12'])
    wb['Parâmetros']['B9'].comment=Comment(spec['analysis']['referencia']['fonte']+'\n'+spec['analysis']['referencia'].get('url',''),'Equipe')
    out=BytesIO();wb.save(out)
    result=BytesIO();ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(BytesIO(out.getvalue())) as src,ZipFile(result,'w',ZIP_DEFLATED) as dest:
        for part in src.infolist():
            data=src.read(part.filename)
            if part.filename.startswith('xl/worksheets/sheet') and part.filename.endswith('.xml'):
                index=int(part.filename.split('sheet')[-1].split('.')[0])-1
                cache={i['address']:i.get('cache') for i in spec['sheets'][index]['cells'] if 'formula' in i}
                root=ET.fromstring(data)
                for c in root.findall('.//s:c',ns):
                    val=cache.get(c.get('r'))
                    if val is not None:
                        if isinstance(val,str):c.set('t','str')
                        v=c.find('s:v',ns)
                        if v is None:v=ET.SubElement(c,'{'+ns['s']+'}v')
                        v.text=str(val)
                data=ET.tostring(root,encoding='utf-8')
            dest.writestr(part,data)
    return result.getvalue()
