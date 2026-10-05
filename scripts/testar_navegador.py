"""Teste da interface simples em navegador REAL (Chromium via Playwright).

Uso (desenvolvimento; Playwright não é dependência do aplicativo):
    pip install playwright && playwright install chromium
    python scripts/testar_navegador.py [pasta_de_capturas]

Sobe o aplicativo numa porta livre, usa as rotas reais e os documentos de
referencias/fontes_salariais_2026_09_07/. A consulta ao TCE-PR é simulada no navegador
(sem internet). Falha com código 1 se qualquer verificação não passar.
"""
import re, sys, tempfile, threading, time
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
SAIDA = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp())
SAIDA.mkdir(parents=True, exist_ok=True)
FONTES = RAIZ / "referencias/fontes_salariais_2026_09_07"
MUNICIPIOS = [{"id": str(i), "nome": n} for i, n in enumerate(["Cafezal do Sul", "Curitiba", "Londrina", "Maringá", "Cruzeiro do Sul"], 1)]
falhas, passos = [], []


def verificar(descricao, condicao, detalhe=""):
    passos.append(descricao)
    print(("  ok   " if condicao else "  FALHA"), descricao, detalhe if not condicao else "")
    if not condicao:
        falhas.append(descricao)


def servidor():
    sys.path.insert(0, str(RAIZ))
    from werkzeug.serving import make_server
    from app import create_app
    srv = make_server("127.0.0.1", 0, create_app(dossie_diretorio=tempfile.mkdtemp()), threaded=True)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def main():
    from playwright.sync_api import sync_playwright
    from openpyxl import load_workbook
    srv = servidor(); base = f"http://127.0.0.1:{srv.server_port}"
    with sync_playwright() as p:
        nav = p.chromium.launch()
        for nome, viewport in (("desktop", {"width": 1280, "height": 900}), ("celular", {"width": 390, "height": 844})):
            print(f"\n== {nome} ==")
            ctx = nav.new_context(viewport=viewport, locale="pt-BR", accept_downloads=True)
            pg = ctx.new_page(); erros_console = []
            pg.on("console", lambda m: erros_console.append(m.text) if m.type == "error" else None)
            pg.on("pageerror", lambda e: erros_console.append(str(e)))
            pg.route("**/api/municipios", lambda r: r.fulfill(json=MUNICIPIOS))
            pg.goto(base + "/"); pg.wait_for_selector("#financeiro-municipio:not([disabled])")
            verificar("tela inicial carrega sem erro no console", not erros_console, str(erros_console))
            verificar("sem rolagem horizontal", pg.evaluate("document.documentElement.scrollWidth<=window.innerWidth+1"))
            sem_rotulo = pg.evaluate("[...document.querySelectorAll('input:not([type=hidden]),select')].filter(e=>!e.labels.length&&!e.getAttribute('aria-label')).map(e=>e.id)")
            verificar("todo campo tem rótulo", not sem_rotulo, str(sem_rotulo))
            pequenos = pg.evaluate("[...document.querySelectorAll('button.principal,#financeiro-municipio,#financeiro-periodo,.abas button')].filter(e=>e.offsetParent!==null&&e.getBoundingClientRect().height<44).map(e=>e.id||e.textContent)")
            verificar("botões e campos principais com pelo menos 44 px de altura", not pequenos, str(pequenos))
            pg.screenshot(path=str(SAIDA / f"{nome}_1_inicio.png"), full_page=True)

            # exemplo pronto + download
            with pg.expect_download() as d:
                pg.click("#exemplo-financeiro")
            arq = SAIDA / f"{nome}_exemplo.xlsx"; d.value.save_as(arq)
            verificar("exemplo financeiro baixa um Excel válido", "Painel financeiro" in load_workbook(arq).sheetnames)
            verificar("resultado mostra o botão de baixar de novo", pg.is_visible("#resultado-financeiro button"))

            # falha da fonte: mensagem em linguagem de usuário e botão liberado
            pg.route("**/api/automacao/financeiro", lambda r: r.fulfill(status=502, json={"erro": "Não foi possível consultar o TCE-PR agora. Tente novamente em alguns minutos.", "tipo": "fonte_externa"}))
            pg.fill("#financeiro-municipio", "Curitiba"); pg.click("#gerar-financeiro")
            pg.wait_for_function("document.querySelector('#financeiro-mensagem').dataset.erro==='true'")
            msg = pg.text_content("#financeiro-mensagem")
            verificar("erro do TCE aparece em português claro", "Tente novamente" in msg and not re.search(r"postback|traceback|500|exception", msg, re.I), msg)
            verificar("botão volta a ficar disponível após o erro", pg.is_enabled("#gerar-financeiro"))
            verificar("barra de andamento some após o erro", not pg.is_visible("#progresso-financeiro"))

            # salários: Londrina
            pg.click("#aba-salarios"); pg.click("#usar-documento")
            pg.set_input_files("#salarios-arquivo", str(FONTES / "londrina.pdf"))
            pg.wait_for_selector("#preparacao-tabela:not([hidden])", timeout=120000)
            rotulos = pg.eval_on_selector_all("#salarios-candidato option", "o=>o.map(x=>x.textContent)")
            verificar("Londrina: rótulos das tabelas são todos distintos", len(set(rotulos)) == len(rotulos) and len(rotulos) == 5, str(rotulos))
            verificar("Londrina: poucas tabelas, sem caixa de filtro", not pg.is_visible("#filtro-candidato"))
            pg.screenshot(path=str(SAIDA / f"{nome}_2_londrina.png"), full_page=True)

            if nome == "desktop":
                # salários: Curitiba (68 tabelas)
                t0 = time.time()
                pg.set_input_files("#salarios-arquivo", str(FONTES / "curitiba.pdf"))
                pg.wait_for_function("document.querySelector('#salarios-mensagem').textContent.includes('Tabela encontrada')", timeout=180000)
                primeira = time.time() - t0
                rotulos = pg.eval_on_selector_all("#salarios-candidato option", "o=>o.map(x=>x.textContent)")
                verificar("Curitiba: 68 tabelas, todas com rótulo distinto", len(rotulos) == 68 and len(set(rotulos)) == 68, f"{len(rotulos)} / {len(set(rotulos))}")
                verificar("Curitiba: agrupadas por grupo ocupacional", pg.eval_on_selector_all("#salarios-candidato optgroup", "g=>g.length") >= 4)
                verificar("Curitiba: caixa de filtro aparece", pg.is_visible("#filtro-candidato"))
                verificar("Curitiba: contagem informada", "68 tabelas" in pg.text_content("#contagem-candidato"))
                pg.screenshot(path=str(SAIDA / "desktop_3_curitiba.png"), full_page=True)
                pg.fill("#filtro-candidato", "M13")
                filtradas = pg.eval_on_selector_all("#salarios-candidato option", "o=>o.map(x=>x.textContent)")
                verificar("filtro 'M13' deixa só as tabelas do Anexo M13", 0 < len(filtradas) <= 3 and all("M13" in x for x in filtradas), str(filtradas))
                verificar("a tabela escolhida aparece descrita por inteiro", "Tabela escolhida:" in pg.text_content("#dados-encontrados") and "M13" in pg.text_content("#dados-encontrados"), pg.text_content("#dados-encontrados"))
                pg.fill("#filtro-candidato", "zzzz")
                verificar("filtro sem resultado orienta o usuário", "Nenhuma tabela" in pg.text_content("#contagem-candidato") and not pg.is_enabled("#gerar-salarios"))
                pg.fill("#filtro-candidato", "Superior")
                verificar("filtro por grupo funciona", pg.eval_on_selector_all("#salarios-candidato option", "o=>o.length") >= 10)
                pg.fill("#filtro-candidato", "S5")
                pg.screenshot(path=str(SAIDA / "desktop_4_filtrado.png"), full_page=True)

                # mesma leitura de novo: instantânea
                t1 = time.time()
                pg.set_input_files("#salarios-arquivo", str(FONTES / "curitiba.pdf"))
                pg.wait_for_function("document.querySelector('#salarios-mensagem').textContent.includes('Tabela encontrada')", timeout=60000)
                segunda = time.time() - t1
                print(f"        1ª leitura {primeira:.1f} s; 2ª leitura do mesmo arquivo {segunda:.1f} s")
                verificar("segunda leitura do mesmo arquivo é rápida (menos de 5 s)", segunda < 5)

                # gerar Excel de uma tabela filtrada
                pg.fill("#filtro-candidato", "M13")
                pg.fill("#complemento-profissao", "Cargo de teste") if pg.is_visible("#complemento-profissao") else None
                pg.select_option("#comparacao-tipo", "percentual"); pg.fill("#comparacao-valor", "5")
                pg.check("#recorte-confirmado")
                with pg.expect_download(timeout=120000) as d:
                    pg.click("#gerar-salarios")
                arq = SAIDA / "curitiba_M13.xlsx"; d.value.save_as(arq)
                wb = load_workbook(arq)
                verificar("Excel salarial gerado a partir da tabela filtrada", "Comparação salarial" in wb.sheetnames)
            ctx.close()
        nav.close()
    srv.shutdown()
    print(f"\n{len(passos) - len(falhas)} de {len(passos)} verificações passaram. Capturas em {SAIDA}")
    sys.exit(1 if falhas else 0)


if __name__ == "__main__":
    main()
