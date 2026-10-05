"""Levanta a estrutura de uma planilha sem interpretar regras ainda não validadas."""

from io import BytesIO

from openpyxl import load_workbook


LIMITE_PLANILHA_BYTES = 15 * 1024 * 1024


def inspecionar_planilha(nome_arquivo, conteudo):
    if not nome_arquivo or not nome_arquivo.lower().endswith(".xlsx"):
        raise ValueError("Envie uma planilha no formato XLSX.")
    if not conteudo:
        raise ValueError("A planilha enviada está vazia.")
    if len(conteudo) > LIMITE_PLANILHA_BYTES:
        raise ValueError("A planilha excede o limite de 15 MB.")

    try:
        workbook = load_workbook(BytesIO(conteudo), data_only=False, read_only=False)
    except (OSError, ValueError, KeyError) as erro:
        raise ValueError("O arquivo não pôde ser interpretado como XLSX.") from erro

    abas = []
    total_formulas = 0
    for planilha in workbook.worksheets:
        formulas = sum(
            1
            for linha in planilha.iter_rows()
            for celula in linha
            if celula.data_type == "f"
        )
        total_formulas += formulas
        abas.append(
            {
                "nome": planilha.title,
                "linhas": planilha.max_row,
                "colunas": planilha.max_column,
                "formulas": formulas,
                "oculta": planilha.sheet_state != "visible",
            }
        )

    return {
        "nome_arquivo": nome_arquivo,
        "tamanho_bytes": len(conteudo),
        "quantidade_abas": len(abas),
        "total_formulas": total_formulas,
        "abas": abas,
        "status": "estrutura_inspecionada_sem_validacao_das_regras",
    }
