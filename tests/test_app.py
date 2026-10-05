import unittest
from io import BytesIO
from tempfile import TemporaryDirectory

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app import create_app
from src.modelos.analise import (
    AnaliseFinanceira,
    AnaliseHistorica,
    DadosFundeb,
    EvolucaoFinanceira,
    LimitesPessoal,
    ResumoHistorico,
)


def analise_falsa(ano=2019):
    return AnaliseFinanceira(
        municipio_id="1490",
        municipio="Curitiba",
        entidade_id="12268",
        entidade="Município de Curitiba",
        ano=ano,
        receita_corrente_liquida=100,
        receita_corrente_liquida_ajustada=100,
        despesa_total_pessoal=40,
        percentual_oficial=40,
        percentual_calculado=40,
        classificacao="normal",
        limites=LimitesPessoal(48.6, 51.3, 54),
        rcl_relatorio_rreo=100,
        diferenca_validacao_rcl=0,
        fonte_receita="RREO",
        fonte_pessoal="RGF",
        coletado_em="2026-08-15T12:00:00-03:00",
        fundeb=DadosFundeb(60, 25, "MDE", "6º Bimestre"),
        evolucao=EvolucaoFinanceira(None, None, None, None),
    )


class ServicoFalso:
    def analisar(
        self,
        municipio_id,
        municipio_nome,
        entidade_id,
        entidade_nome,
        ano,
        incluir_fundeb=True,
    ):
        return analise_falsa(ano)

    def analisar_historico(
        self,
        municipio_id,
        municipio_nome,
        entidade_id,
        entidade_nome,
        ano_inicial,
        ano_final,
        incluir_fundeb=True,
    ):
        if ano_final < ano_inicial:
            raise ValueError("O exercício final deve ser maior ou igual ao inicial.")
        resultado = analise_falsa(ano_inicial)
        return AnaliseHistorica(
            municipio_id=municipio_id,
            municipio=municipio_nome,
            entidade_id=entidade_id,
            entidade=entidade_nome,
            ano_inicial=ano_inicial,
            ano_final=ano_final,
            resultados=(resultado,),
            erros=(),
            resumo=ResumoHistorico(1, 0, 0, 0, 0, 40, 40, ano_inicial),
            coletado_em="2026-08-15T12:00:00-03:00",
        )


class AppTestCase(unittest.TestCase):
    def setUp(self):
        self.diretorio_temporario = TemporaryDirectory()
        self.addCleanup(self.diretorio_temporario.cleanup)
        app = create_app(ServicoFalso(), dossie_diretorio=self.diretorio_temporario.name)
        app.config.update(TESTING=True)
        self.client = app.test_client()

    def test_pagina_inicial(self):
        resposta = self.client.get("/")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("Análise Municipal".encode(), resposta.data)

    def test_interface_separa_consulta_real_da_demo(self):
        resposta = self.client.get("/avancado")
        self.assertIn(b"MVP 0.9.0", resposta.data)
        self.assertIn("Curitiba é a demonstração, não o limite".encode(), resposta.data)
        self.assertIn(b"Consulta real:", resposta.data)
        self.assertIn(b"Demo offline:", resposta.data)

    def test_amostra_anual_de_curitiba_nao_fica_presa_em_2019(self):
        for ano in (2019, 2025):
            with self.subTest(ano=ano):
                resposta = self.client.get(f"/api/amostras/curitiba/{ano}")
                self.assertEqual(resposta.status_code, 200)
                self.assertEqual(resposta.get_json()["ano"], ano)
                self.assertEqual(
                    resposta.get_json()["origem_dados"],
                    "Amostra preservada de Curitiba",
                )

    def test_amostra_anual_informa_ano_nao_preservado(self):
        resposta = self.client.get("/api/amostras/curitiba/2018")
        self.assertEqual(resposta.status_code, 404)
        self.assertEqual(resposta.get_json()["anos_disponiveis"], list(range(2019, 2026)))

    def test_amostra_historica_informa_origem(self):
        resposta = self.client.get("/api/amostras/curitiba-historico")
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["ano_inicial"], 2019)
        self.assertEqual(resposta.get_json()["ano_final"], 2025)
        self.assertEqual(
            resposta.get_json()["origem_dados"],
            "Amostra preservada de Curitiba",
        )

    def test_analise_valida(self):
        resposta = self.client.post(
            "/api/analisar",
            json={
                "municipio_id": "1490",
                "municipio_nome": "Curitiba",
                "entidade_id": "12268",
                "entidade_nome": "Município de Curitiba",
                "ano": 2019,
            },
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["classificacao"], "normal")
        self.assertEqual(resposta.get_json()["fundeb"]["receitas_recebidas"], 60)

    def test_analise_historica_valida(self):
        resposta = self.client.post(
            "/api/analisar-historico",
            json={
                "municipio_id": "1490",
                "municipio_nome": "Curitiba",
                "entidade_id": "12268",
                "entidade_nome": "Município de Curitiba",
                "ano_inicial": 2019,
                "ano_final": 2025,
            },
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["resumo"]["anos_analisados"], 1)

    def test_campos_obrigatorios(self):
        resposta = self.client.post("/api/analisar", json={"ano": 2019})
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("Campos obrigatórios", resposta.get_json()["erro"])

    def test_inspeciona_pdf_valido(self):
        escritor = PdfWriter()
        escritor.add_blank_page(width=100, height=100)
        conteudo = BytesIO()
        escritor.write(conteudo)
        resposta = self.client.post(
            "/api/salarios/inspecionar-pdf",
            data={"arquivo": (BytesIO(conteudo.getvalue()), "tabela.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["quantidade_paginas"], 1)
        self.assertTrue(resposta.get_json()["provavelmente_escaneado"])

    def test_inspeciona_planilha_valida(self):
        workbook = Workbook()
        workbook.active["A1"] = "=1+1"
        conteudo = BytesIO()
        workbook.save(conteudo)
        resposta = self.client.post(
            "/api/materiais/inspecionar-planilha",
            data={"arquivo": (BytesIO(conteudo.getvalue()), "original.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["total_formulas"], 1)

    def test_amostra_e_analise_salarial(self):
        amostra = self.client.get("/api/salarios/amostras/cafezal-do-sul-2026")
        self.assertEqual(amostra.status_code, 200)
        self.assertEqual(amostra.get_json()["municipio"], "Cafezal do Sul")
        resposta = self.client.post("/api/salarios/analisar", json=amostra.get_json())
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["resumo"]["celulas_abaixo_referencia"], 0)

    def test_lista_fontes_salariais_sem_limitar_a_cafezal(self):
        resposta = self.client.get("/api/salarios/fontes")
        self.assertEqual(resposta.status_code, 200)
        municipios = {item["municipio"] for item in resposta.get_json()}
        self.assertEqual(len(municipios), 6)
        self.assertTrue({"Cafezal do Sul", "Curitiba", "Londrina"}.issubset(municipios))

    def test_registra_dossie_salarial_sem_executar_comparacao(self):
        conteudo = BytesIO()
        documento = Document()
        documento.add_paragraph("Tabela salarial do magistério - referência 2026")
        documento.save(conteudo)
        conteudo.seek(0)
        resposta = self.client.post(
            "/api/salarios/dossies",
            data={
                "municipio": "Londrina",
                "uf": "PR",
                "ano_referencia": "2026",
                "fonte_url": "https://portal.londrina.pr.gov.br/pessoal",
                "tipo_fonte": "portal_prefeitura",
                "observacoes": "Coleta de teste.",
                "tabela_vigente": (conteudo, "tabela.docx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(resposta.status_code, 201)
        dados = resposta.get_json()
        self.assertEqual(dados["municipio"], "Londrina")
        self.assertEqual(dados["status"], "parcial_para_triagem")
        self.assertNotIn("tabela_referencia", dados)

    def test_exporta_analise_salarial_em_excel_e_pdf(self):
        amostra = self.client.get("/api/salarios/amostras/cafezal-do-sul-2026").get_json()
        analise = self.client.post("/api/salarios/analisar", json=amostra).get_json()
        excel = self.client.post("/api/salarios/exportar-excel", json=analise)
        pdf = self.client.post("/api/salarios/relatorio-pdf", json=analise)
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(excel.data.startswith(b"PK"))
        self.assertTrue(pdf.data.startswith(b"%PDF"))

    def test_importa_tabela_salarial_docx(self):
        documento = Document()
        documento.add_paragraph("Anexo A da LC 064/2026")
        documento.add_paragraph("Nível A - Magistério 2.565,31 2.616,61")
        documento.add_paragraph("Nível B - Licenciatura Plena 2.873,14 2.930,60")
        documento.add_paragraph("Percentual entre classes = 2,00%")
        documento.add_paragraph("Nível B = Nível A acrescido de 12,00%")
        conteudo = BytesIO()
        documento.save(conteudo)
        resposta = self.client.post(
            "/api/salarios/importar-documento",
            data={
                "arquivo": (BytesIO(conteudo.getvalue()), "tabela.docx"),
                "municipio": "Cafezal do Sul",
                "uf": "PR",
                "ano": "2026",
                "piso_nacional_40h": "5130.63",
                "jornada_semanal": "20",
                "modo_arredondamento": "truncar",
                "fonte_piso": "Lei Federal nº 15.437/2026",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(len(resposta.get_json()["tabela_atual"]), 2)
        self.assertEqual(resposta.get_json()["parametros"]["progressao_classes_percentual"], 2)

    def test_importa_romanos_e_intersticio_pela_api(self):
        documento = Document()
        documento.add_paragraph("Anexo Único da LC 010/2024")
        documento.add_paragraph("Nível I - Magistério 2.400,00 2.520,00")
        documento.add_paragraph("Nível II - Licenciatura 2.640,00 2.772,00")
        documento.add_paragraph("Interstício de 5% entre classes")
        documento.add_paragraph("Nível II corresponde ao Nível I acrescido de 10%")
        conteudo = BytesIO()
        documento.save(conteudo)
        resposta = self.client.post(
            "/api/salarios/importar-documento",
            data={
                "arquivo": (BytesIO(conteudo.getvalue()), "romanos.docx"),
                "municipio": "Município Exemplo",
                "ano": "2026",
                "piso_nacional_40h": "4800",
                "jornada_semanal": "20",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(resposta.status_code, 200)
        dados = resposta.get_json()
        self.assertEqual([nivel["codigo"] for nivel in dados["tabela_atual"]], ["I", "II"])
        self.assertEqual(dados["parametros"]["progressoes_classes_percentuais"], [5])

    def test_api_rejeita_perda_silenciosa_de_nivel(self):
        documento = Document()
        documento.add_paragraph("Nível A - Magistério 2.400,00 2.448,00")
        documento.add_paragraph("Nível @ Licenciatura 2.600,00 2.652,00")
        documento.add_paragraph("Percentual entre classes = 2,00%")
        conteudo = BytesIO()
        documento.save(conteudo)
        resposta = self.client.post(
            "/api/salarios/importar-documento",
            data={
                "arquivo": (BytesIO(conteudo.getvalue()), "invalido.docx"),
                "municipio": "Município Exemplo",
                "ano": "2026",
                "piso_nacional_40h": "4800",
                "jornada_semanal": "20",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("não puderam ser interpretadas", resposta.get_json()["erro"])


if __name__ == "__main__":
    unittest.main()
