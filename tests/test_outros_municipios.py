"""Regressões do QA exploratório com estruturas diferentes de Cafezal do Sul."""

import unittest
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from src.exportacao.excel_salarial import gerar_excel_salarial
from src.exportacao.relatorio_pdf import gerar_relatorio_salarial_pdf
from src.salarial.analise_service import analisar_tabela_salarial
from src.salarial.documento_service import extrair_tabela_salarial_texto


TEXTO_ROMANOS = """
Anexo Único da LC 010/2024
Nível I - Ensino Médio Magistério 2.400,00 2.448,00 2.496,96
Nível II - Licenciatura Curta 2.640,00 2.692,80 2.746,66
Nível III - Licenciatura Plena 3.000,00 3.060,00 3.121,20
Percentual entre classes = 2,00%
Nível II = Nível I acrescido de 10,00%
Nível III = Nível I acrescido de 25,00%
"""

TEXTO_NUMERICO = """
Nível 01 - Professor 2400,00 2448,00
Nível 02 - Professor Pleno 2600,00 2652,00
Percentual entre classes = 2,00%
Nível 02 = Nível 01 acrescido de R$ 200,00
"""

TEXTO_INTERSTICIO = """
Lei Municipal 1.234/2023 - Plano de Carreira do Magistério
Nível A - Professor 2.400,00 2.520,00 2.646,00
Nível B - Professor Licenciado 2.640,00 2.772,00 2.910,60
Interstício de 5% entre classes.
Nível B corresponde ao Nível A acrescido de 10%.
"""


class OutrosMunicipiosTestCase(unittest.TestCase):
    def test_importa_todos_os_niveis_romanos(self):
        resultado = extrair_tabela_salarial_texto(TEXTO_ROMANOS, "romanos.docx")
        self.assertEqual([nivel["codigo"] for nivel in resultado["niveis"]], ["I", "II", "III"])
        self.assertEqual(resultado["niveis"][1]["regra_vertical"]["nivel_referencia"], "I")
        self.assertEqual(resultado["niveis"][2]["acrescimo_percentual"], 25)

    def test_importa_niveis_numericos_sem_perder_milhar(self):
        resultado = extrair_tabela_salarial_texto(TEXTO_NUMERICO, "numericos.docx")
        self.assertEqual([nivel["codigo"] for nivel in resultado["niveis"]], ["01", "02"])
        self.assertEqual(resultado["niveis"][0]["valores"][0], 2400)
        self.assertEqual(resultado["niveis"][1]["regra_vertical"]["tipo"], "valor_fixo")
        self.assertEqual(resultado["niveis"][1]["acrescimo_valor_fixo"], 200)

    def test_reconhece_intersticio_como_progressao(self):
        resultado = extrair_tabela_salarial_texto(TEXTO_INTERSTICIO, "intersticio.docx")
        self.assertEqual(resultado["progressoes_classes_percentuais"], [5, 5])
        self.assertEqual(resultado["niveis"][1]["acrescimo_percentual"], 10)

    def test_nao_descarta_linha_salarial_desconhecida_silenciosamente(self):
        texto = """
Nível A - Professor 2.400,00 2.448,00
Nível @ Professor Pleno 2.600,00 2.652,00
Percentual entre classes = 2,00%
"""
        with self.assertRaisesRegex(ValueError, "não puderam ser interpretadas"):
            extrair_tabela_salarial_texto(texto)

    def test_progressao_nao_uniforme_valida_cada_transicao(self):
        dados = {
            "municipio": "Município Exemplo B",
            "ano": 2026,
            "jornada_semanal": 20,
            "piso_nacional_40h": "5130.63",
            "progressoes_classes_percentuais": ["5", "5", "10", "10"],
            "modo_arredondamento": "truncar",
            "niveis": [
                {
                    "codigo": "A",
                    "valores": ["2565.31", "2693.57", "2828.24", "3111.06", "3422.16"],
                }
            ],
        }
        resultado = analisar_tabela_salarial(dados)
        self.assertEqual(resultado["resumo"]["divergencias_estrutura"], 0)
        self.assertFalse(resultado["parametros"]["progressao_classes_uniforme"])
        self.assertEqual(
            resultado["parametros"]["progressoes_classes_percentuais"],
            [5, 5, 10, 10],
        )

    def test_acrescimo_vertical_fixo_em_reais(self):
        dados = {
            "municipio": "Município Exemplo C",
            "ano": 2026,
            "jornada_semanal": 20,
            "piso_nacional_40h": "5130.63",
            "progressao_classes_percentual": "2",
            "modo_arredondamento": "truncar",
            "niveis": [
                {
                    "codigo": "A",
                    "valores": ["2565.31", "2616.61", "2668.94"],
                },
                {
                    "codigo": "B",
                    "regra_vertical": {
                        "tipo": "valor_fixo",
                        "valor": "300.00",
                        "nivel_referencia": "A",
                    },
                    "valores": ["2865.31", "2922.61", "2981.06"],
                },
            ],
        }
        resultado = analisar_tabela_salarial(dados)
        self.assertEqual(resultado["resumo"]["divergencias_estrutura"], 0)
        self.assertEqual(resultado["tabela_referencia"][1]["valores"], [2865.31, 2922.61, 2981.06])
        self.assertEqual(resultado["tabela_atual"][1]["regra_vertical"]["tipo"], "valor_fixo")

        workbook = load_workbook(
            BytesIO(gerar_excel_salarial(resultado)), data_only=False
        )
        self.assertEqual(workbook["Parâmetros"]["C20"].value, "valor_fixo")
        self.assertEqual(workbook["Parâmetros"]["D20"].value, 300)
        self.assertIn("+'Parâmetros'!D20", workbook["Comparação salarial"]["D23"].value)

        pdf = gerar_relatorio_salarial_pdf(resultado)
        texto = "\n".join(
            pagina.extract_text() or "" for pagina in PdfReader(BytesIO(pdf)).pages
        )
        self.assertIn("valor fixo", texto.lower())

    def test_excel_usa_percentual_de_cada_transicao(self):
        dados = {
            "municipio": "Município Exemplo D",
            "jornada_semanal": 20,
            "piso_nacional_40h": "5130.63",
            "progressoes_classes_percentuais": [5, 5, 10, 10],
            "niveis": [
                {
                    "codigo": "A",
                    "valores": [2565.31, 2693.57, 2828.24, 3111.06, 3422.16],
                }
            ],
        }
        resultado = analisar_tabela_salarial(dados)
        workbook = load_workbook(
            BytesIO(gerar_excel_salarial(resultado)), data_only=False
        )
        parametros = workbook["Parâmetros"]
        self.assertEqual(parametros["C24"].value, 5)
        self.assertEqual(parametros["C26"].value, 10)
        referencia = workbook["Comparação salarial"]
        self.assertIn("C24", referencia["E20"].value)
        self.assertIn("C26", referencia["G20"].value)

    def test_rejeita_quantidade_errada_de_progressoes(self):
        dados = {
            "jornada_semanal": 20,
            "piso_nacional_40h": 5000,
            "progressoes_classes_percentuais": [2],
            "niveis": [{"codigo": "A", "valores": [2500, 2550, 2601]}],
        }
        with self.assertRaisesRegex(ValueError, "exatamente 2"):
            analisar_tabela_salarial(dados)

    def test_percentual_em_nota_nao_vira_classe_salarial(self):
        texto = """
Nível A - Professor 2.400,00 2.448,00 (reajuste anual de 5,00%)
Nível B - Licenciatura 2.640,00 2.692,80 (reajuste anual de 5,00%)
Percentual entre classes = 2,00%
Nível B = Nível A acrescido de 10,00%
"""
        resultado = extrair_tabela_salarial_texto(texto)
        self.assertEqual(resultado["classes"], [1, 2])
        self.assertEqual(resultado["niveis"][0]["valores"], [2400, 2448])
        self.assertEqual(resultado["niveis"][1]["valores"], [2640, 2692.8])

    def test_rejeita_regras_verticais_conflitantes(self):
        texto = """
Nível A - Professor 2.400,00 2.448,00
Nível B - Licenciatura 2.640,00 2.692,80
Percentual entre classes = 2,00%
Nível B = Nível A acrescido de 12,00%
Em caso de dúvida, considera-se que o Nível B = Nível A acrescido de 20,00%
"""
        with self.assertRaisesRegex(ValueError, "regras verticais conflitantes"):
            extrair_tabela_salarial_texto(texto)

    def test_rejeita_declaracoes_gerais_de_progressao_conflitantes(self):
        texto = """
Nível A - Professor 2.400,00 2.520,00 2.646,00
Nível B - Licenciatura 2.640,00 2.772,00 2.910,60
Percentual entre classes = 5,00%
Nota de rodapé: revogado o artigo anterior, o percentual entre classes passa a ser de 8,00%
Nível B = Nível A acrescido de 10,00%
"""
        with self.assertRaisesRegex(ValueError, "declarações conflitantes"):
            extrair_tabela_salarial_texto(texto)


if __name__ == "__main__":
    unittest.main()
