import json
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from pypdf import PdfWriter

from src.persistencia.dossie_salarial import DossieSalarialStore
from src.salarial.aquisicao_service import (
    inspecionar_documento_bruto,
    registrar_dossie_salarial,
)
from src.salarial.fontes_service import listar_fontes_salariais


def docx_textual():
    documento = Document()
    documento.add_heading("Tabela salarial do magistério", level=1)
    documento.add_paragraph("Nível A - Classe 1 - R$ 2.565,31")
    conteudo = BytesIO()
    documento.save(conteudo)
    return conteudo.getvalue()


class AquisicaoSalarialTestCase(unittest.TestCase):
    def test_catalogo_tem_seis_fontes_oficiais_sem_duplicidade(self):
        fontes = listar_fontes_salariais()
        self.assertEqual(len(fontes), 6)
        self.assertEqual(len({item["id"] for item in fontes}), 6)
        self.assertIn("Cafezal do Sul", {item["municipio"] for item in fontes})
        self.assertIn("Curitiba", {item["municipio"] for item in fontes})
        self.assertTrue(all(item["pagina_oficial"].startswith("https://") for item in fontes))
        self.assertTrue(all(item["coleta_automatica"] is False for item in fontes))

    def test_pdf_sem_texto_e_registrado_com_alerta_de_ocr(self):
        escritor = PdfWriter()
        escritor.add_blank_page(width=100, height=100)
        conteudo = BytesIO()
        escritor.write(conteudo)
        resultado = inspecionar_documento_bruto(
            "tabela.pdf", conteudo.getvalue(), "tabela_vigente"
        )
        self.assertTrue(resultado["requer_ocr"])
        self.assertFalse(resultado["texto_extraivel"])
        self.assertEqual(resultado["quantidade_paginas"], 1)

    def test_docx_textual_fica_pronto_para_triagem(self):
        resultado = inspecionar_documento_bruto(
            "tabela.docx", docx_textual(), "tabela_vigente"
        )
        self.assertTrue(resultado["texto_extraivel"])
        self.assertGreater(resultado["quantidade_caracteres"], 20)
        self.assertFalse(resultado["requer_conversao"])

    def test_doc_legado_e_preservado_sem_exigir_libreoffice_na_aquisicao(self):
        resultado = inspecionar_documento_bruto(
            "tabela.doc", b"conteudo legado para preservar", "tabela_vigente"
        )
        self.assertTrue(resultado["requer_conversao"])
        self.assertEqual(resultado["status_inspecao"], "registrado_aguardando_conversao")

    def test_registra_originais_manifesto_hash_e_pendencias(self):
        with TemporaryDirectory() as temporario:
            raiz = Path(temporario) / "salarios"
            tabela = docx_textual()
            resultado = registrar_dossie_salarial(
                DossieSalarialStore(raiz),
                {
                    "municipio": "Londrina",
                    "uf": "PR",
                    "ano": "2026",
                    "fonte_url": "https://portal.londrina.pr.gov.br/pessoal",
                    "tipo_fonte": "portal_prefeitura",
                    "observacoes": "Conferir a revisão anual.",
                },
                [
                    {
                        "categoria": "tabela_vigente",
                        "nome_original": "Anexo III.docx",
                        "conteudo": tabela,
                    }
                ],
            )
            self.assertEqual(resultado["status"], "parcial_para_triagem")
            self.assertEqual(
                resultado["itens_recomendados_ausentes"],
                ["plano_carreira", "ato_reajuste"],
            )
            manifestos = list(raiz.rglob("manifesto.json"))
            self.assertEqual(len(manifestos), 1)
            salvo = json.loads(manifestos[0].read_text(encoding="utf-8"))
            self.assertEqual(salvo["arquivos"][0]["sha256"], resultado["arquivos"][0]["sha256"])
            original = manifestos[0].parent / salvo["arquivos"][0]["nome_armazenado"]
            self.assertEqual(original.read_bytes(), tabela)

    def test_bloqueia_dossie_sem_tabela_vigente(self):
        with TemporaryDirectory() as temporario:
            with self.assertRaisesRegex(ValueError, "tabela salarial vigente"):
                registrar_dossie_salarial(
                    DossieSalarialStore(temporario),
                    {
                        "municipio": "Maringá",
                        "uf": "PR",
                        "ano": 2026,
                        "fonte_url": "https://transparencia.maringa.pr.gov.br/",
                        "tipo_fonte": "portal_prefeitura",
                    },
                    [],
                )


if __name__ == "__main__":
    unittest.main()
