import unittest,shutil
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image,ImageDraw,ImageFont
from src.salarial.revisao_service import extrair_para_revisao

class FormatosTest(unittest.TestCase):
 def validar(self,nome,b):
  d=extrair_para_revisao(nome,b)['candidatos'];self.assertEqual(len(d),1)
  self.assertEqual(d[0]['niveis'][1]['valores'],[1500,1575]);return d
 def test_csv_e_tsv(self):
  for sep in [';', '\t']:
   self.validar('t.csv', ('Nivel'+sep+'A'+sep+'B\nI'+sep+'1200,00'+sep+'1260,00\nII'+sep+'1500,00'+sep+'1575,00').encode())
 def test_xlsx_cabecalho_apos_titulo(self):
  w=Workbook();s=w.active;s.append(['Tabela da carreira']);s.append(['Nível','A','B']);s.append(['I',1200,1260]);s.append(['II',1500,1575]);b=BytesIO();w.save(b);self.validar('t.xlsx',b.getvalue())
 def test_docx_tabela_nativa(self):
  w=Document();w.add_paragraph('Município: Cidade Nova');t=w.add_table(rows=3,cols=3)
  for row,vs in zip(t.rows,[['Nível','A','B'],['I','1200,00','1260,00'],['II','1500,00','1575,00']]):
   for c,v in zip(row.cells,vs):c.text=v
  b=BytesIO();w.save(b);self.validar('t.docx',b.getvalue())
 def test_html_matriz_nao_municipal(self):
  self.validar('t.html',b'<table><tr><th>Nivel</th><th>A</th><th>B</th></tr><tr><td>I</td><td>1200,00</td><td>1260,00</td></tr><tr><td>II</td><td>1500,00</td><td>1575,00</td></tr></table>')
 def test_pdf_textual_sem_nome_municipal(self):
  b=BytesIO();c=canvas.Canvas(b)
  for y,t in [(750,'Nivel A B'),(720,'I 1200,00 1260,00'),(690,'II 1500,00 1575,00')]:c.drawString(50,y,t)
  c.save();self.validar('t.pdf',b.getvalue())
 @unittest.skipUnless(shutil.which('tesseract'),'OCR opcional não instalado')
 def test_pdf_imagem_ocr(self):
  im=Image.new('RGB',(1200,450),'white');dr=ImageDraw.Draw(im);font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',35)
  for y,vs in [(40,['Nivel','A','B']),(130,['I','1200,00','1260,00']),(220,['II','1500,00','1575,00'])]:
   for x,v in zip([70,450,850],vs):dr.text((x,y),v,font=font,fill='black')
  b=BytesIO();c=canvas.Canvas(b,pagesize=(600,225));c.drawImage(ImageReader(im),0,0,width=600,height=225);c.save()
  d=self.validar('imagem.pdf',b.getvalue());self.assertTrue(d[0]['origem_ocr'])
 def test_valor_formula_sem_cache_rejeitado(self):
  w=Workbook();s=w.active;s.append(['Nível','A','B']);s.append(['I',1200,'=1200*1.05']);b=BytesIO();w.save(b)
  r=extrair_para_revisao('t.xlsx',b.getvalue());self.assertFalse(r['candidatos']);self.assertTrue(r['pendencias'])

 def test_xls_legado(self):
  p=Path(__file__).resolve().parents[1]/'referencias/testes_leitor_geral/matriz.xls'
  self.validar('t.xls',p.read_bytes())
 @unittest.skipUnless(shutil.which('soffice') or shutil.which('antiword'),'conversor DOC opcional não instalado')
 def test_doc_real_do_professor(self):
  p=Path(__file__).resolve().parents[1]/'referencias/materiais_professor/Anexo_A_LC_064_2026_Magisterio.doc'
  r=extrair_para_revisao(p.name,p.read_bytes());self.assertEqual(len(r['candidatos']),1)
  self.assertEqual(len(r['candidatos'][0]['niveis']),3)
 def test_ocr_ausente_devolve_pendencia(self):
  from unittest.mock import patch
  b=BytesIO();c=canvas.Canvas(b);c.showPage();c.save()
  with patch('src.salarial.formatos_gerais.shutil.which',return_value=None):
   r=extrair_para_revisao('imagem.pdf',b.getvalue())
  self.assertFalse(r['candidatos']);self.assertTrue(any('Tesseract' in p for p in r['pendencias']))
