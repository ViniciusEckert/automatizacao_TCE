"""Exportação das análises anual e histórica para Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from src.exportacao.cache_xlsx import injetar_cache
from src.exportacao.painel_professor import criar_painel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CORES_STATUS = {
    "normal": "C7DB78",
    "alerta": "E6BA59",
    "prudencial": "D98D42",
    "limite_excedido": "BD5C4C",
}
VERDE_ESCURO = "124B3A"
VERDE = "1E6B53"
VERDE_CLARO = "EAF4EC"
BRANCO = "FFFFFF"
CINZA = "DCD8CC"
MOEDA = 'R$ #,##0.00'
PERCENTUAL = "0.00%"


def _sem_suavizar(grafico):
    """Linhas retas entre os anos: a suavização do openpyxl inventa picos que não existem nos dados."""
    for serie in grafico.series:
        serie.smooth = False
    return grafico


def _titulo(planilha, texto, ultima_coluna):
    planilha.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_coluna)
    celula = planilha.cell(1, 1, texto)
    celula.font = Font(size=16, bold=True, color=BRANCO)
    celula.fill = PatternFill("solid", fgColor=VERDE_ESCURO)
    celula.alignment = Alignment(horizontal="center")
    planilha.row_dimensions[1].height = 28


def _cabecalho(linha):
    for celula in linha:
        celula.font = Font(bold=True, color=BRANCO)
        celula.fill = PatternFill("solid", fgColor=VERDE)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _finalizar(workbook):
    arquivo = BytesIO()
    workbook.save(arquivo)
    return arquivo.getvalue()


def _configurar_impressao(planilha, area, paisagem=True, altura=1):
    planilha.print_area = area
    planilha.page_setup.orientation = "landscape" if paisagem else "portrait"
    planilha.page_setup.paperSize = planilha.PAPERSIZE_A4
    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = altura
    planilha.sheet_properties.pageSetUpPr.fitToPage = True
    planilha.page_margins.left = 0.25
    planilha.page_margins.right = 0.25
    planilha.page_margins.top = 0.35
    planilha.page_margins.bottom = 0.35


def _validar_anual(dados):
    obrigatorios = (
        "municipio",
        "entidade",
        "ano",
        "receita_corrente_liquida",
        "receita_corrente_liquida_ajustada",
        "despesa_total_pessoal",
        "percentual_oficial",
        "classificacao",
    )
    ausentes = [campo for campo in obrigatorios if campo not in dados]
    if ausentes:
        raise ValueError(f"Dados insuficientes para exportar: {', '.join(ausentes)}")


def gerar_excel_anual(dados):
    _validar_anual(dados)
    workbook = Workbook()
    resumo = workbook.active
    resumo.title = "Resumo"
    resumo.sheet_view.showGridLines = False
    _titulo(resumo, "ANÁLISE FINANCEIRA DO MAGISTÉRIO MUNICIPAL", 2)

    fundeb = dados.get("fundeb") or {}
    linhas = [
        ("Município", dados["municipio"]),
        ("Entidade", dados["entidade"]),
        ("Exercício", dados["ano"]),
        ("Receita Corrente Líquida", dados["receita_corrente_liquida"]),
        ("Receita Corrente Líquida Ajustada", dados["receita_corrente_liquida_ajustada"]),
        ("Despesa Total com Pessoal", dados["despesa_total_pessoal"]),
        ("Percentual oficial", dados["percentual_oficial"] / 100),
        ("Percentual recalculado", dados.get("percentual_calculado", 0) / 100),
        ("Classificação", dados["classificacao"]),
        ("Receitas recebidas do FUNDEB", fundeb.get("receitas_recebidas")),
        ("Resultado líquido das transferências do FUNDEB", fundeb.get("resultado_liquido_transferencias")),
        ("Aplicação constitucional em MDE", fundeb.get("percentual_aplicacao_mde")),
        ("Coletado em", dados.get("coletado_em", "Não informado")),
    ]
    for linha in linhas:
        resumo.append(linha)

    borda = Border(bottom=Side(style="thin", color=CINZA))
    for linha in resumo.iter_rows(min_row=2, max_row=resumo.max_row, min_col=1, max_col=2):
        linha[0].font = Font(bold=True, color="17322C")
        linha[0].border = borda
        linha[1].border = borda
    for linha in range(5, 8):
        resumo.cell(linha, 2).number_format = MOEDA
    resumo["B8"].number_format = PERCENTUAL
    resumo["B9"].number_format = PERCENTUAL
    resumo["B10"].fill = PatternFill(
        "solid", fgColor=CORES_STATUS.get(dados["classificacao"], CINZA)
    )
    resumo["B11"].number_format = MOEDA
    resumo["B12"].number_format = MOEDA
    resumo["B13"].number_format = PERCENTUAL
    if resumo["B13"].value is not None:
        resumo["B13"].value = resumo["B13"].value / 100
    resumo.column_dimensions["A"].width = 46
    resumo.column_dimensions["B"].width = 34

    grafico = BarChart()
    grafico.title = "RCL ajustada x Despesa com Pessoal"
    grafico.y_axis.title = "Valor em R$"
    grafico.add_data(Reference(resumo, min_col=2, min_row=6, max_row=7), titles_from_data=False)
    grafico.set_categories(Reference(resumo, min_col=1, min_row=6, max_row=7))
    grafico.height = 7
    grafico.width = 13
    resumo.add_chart(grafico, "D3")

    fontes = workbook.create_sheet("Fontes e validação")
    fontes.sheet_view.showGridLines = False
    fontes.append(["Campo", "Valor"])
    fontes.append(["Relatório de Receita", dados.get("fonte_receita", "Não informado")])
    fontes.append(["Relatório de Pessoal", dados.get("fonte_pessoal", "Não informado")])
    fontes.append(["Relatório do FUNDEB", fundeb.get("fonte", "Não disponível")])
    fontes.append(["Período do FUNDEB", fundeb.get("periodo", "Não disponível")])
    fontes.append(["RCL no RREO", dados.get("rcl_relatorio_rreo")])
    fontes.append(["Diferença de validação da RCL", dados.get("diferenca_validacao_rcl")])
    fontes.append(["ID do município", dados.get("municipio_id")])
    fontes.append(["ID da entidade", dados.get("entidade_id")])
    fontes.append(["Avisos", " | ".join(dados.get("avisos") or []) or "Nenhum"])
    fontes.column_dimensions["A"].width = 34
    fontes.column_dimensions["B"].width = 85
    _cabecalho(fontes[1])
    for celula in fontes["B"]:
        celula.alignment = Alignment(wrap_text=True, vertical="top")

    _configurar_impressao(resumo, f"A1:P{resumo.max_row}", altura=2)
    _configurar_impressao(fontes, f"A1:B{fontes.max_row}", paisagem=False)

    return _finalizar(workbook)



def gerar_excel_historico(dados):
    resultados = dados.get("resultados") or []
    if not resultados:
        raise ValueError("A análise histórica não possui exercícios para exportar.")

    workbook = Workbook()
    resumo = workbook.active
    resumo.title = "Resumo histórico"
    resumo.sheet_view.showGridLines = False
    _titulo(resumo, "HISTÓRICO FINANCEIRO DO MAGISTÉRIO MUNICIPAL", 4)
    resumo.append(["Município", dados.get("municipio"), "Entidade", dados.get("entidade")])
    resumo.append(["Período", f"{dados.get('ano_inicial')}–{dados.get('ano_final')}", "Coletado em", dados.get("coletado_em")])
    resumo.append([])
    resumo.append(["Indicador", "Resultado"])
    resumo_dados = dados.get("resumo") or {}
    indicadores = [
        ("Exercícios analisados", resumo_dados.get("anos_analisados")),
        ("Exercícios com erro", resumo_dados.get("anos_com_erro")),
        ("Evolução acumulada da RCL ajustada", resumo_dados.get("evolucao_acumulada_rcl_ajustada")),
        ("Evolução acumulada da despesa", resumo_dados.get("evolucao_acumulada_despesa_pessoal")),
        ("Evolução acumulada das receitas do FUNDEB", resumo_dados.get("evolucao_acumulada_receitas_fundeb")),
        ("Média de comprometimento", resumo_dados.get("media_comprometimento")),
        ("Maior comprometimento", resumo_dados.get("maior_comprometimento")),
        ("Ano do maior comprometimento", resumo_dados.get("ano_maior_comprometimento")),
    ]
    for indicador, valor in indicadores:
        resumo.append([indicador, valor])
    _cabecalho(resumo[5])
    for linha in (8, 9, 10, 11, 12):
        resumo.cell(linha, 2).number_format = PERCENTUAL
        if resumo.cell(linha, 2).value is not None:
            resumo.cell(linha, 2).value = resumo.cell(linha, 2).value / 100
    resumo.column_dimensions["A"].width = 46
    resumo.column_dimensions["B"].width = 28
    resumo.column_dimensions["C"].width = 18
    resumo.column_dimensions["D"].width = 44

    historico = workbook.create_sheet("Histórico")
    historico.sheet_view.showGridLines = False
    cabecalhos = [
        "Ano",
        "RCL",
        "RCL ajustada",
        "Despesa com pessoal",
        "% oficial",
        "% recalculado",
        "Classificação",
        "Receitas FUNDEB",
        "Resultado líquido FUNDEB",
        "Aplicação MDE",
        "Δ RCL ajustada",
        "Δ despesa",
        "Δ FUNDEB",
        "Δ comprometimento (p.p.)",
        "Avisos",
        "Limite de alerta",
        "Limite prudencial",
        "Limite máximo",
    ]
    historico.append(cabecalhos)
    for item in resultados:
        fundeb = item.get("fundeb") or {}
        evolucao = item.get("evolucao") or {}
        historico.append(
            [
                item.get("ano"),
                item.get("receita_corrente_liquida"),
                item.get("receita_corrente_liquida_ajustada"),
                item.get("despesa_total_pessoal"),
                item.get("percentual_oficial", 0) / 100,
                item.get("percentual_calculado", 0) / 100,
                item.get("classificacao"),
                fundeb.get("receitas_recebidas"),
                fundeb.get("resultado_liquido_transferencias"),
                None if fundeb.get("percentual_aplicacao_mde") is None else fundeb["percentual_aplicacao_mde"] / 100,
                None if evolucao.get("rcl_ajustada_percentual") is None else evolucao["rcl_ajustada_percentual"] / 100,
                None if evolucao.get("despesa_pessoal_percentual") is None else evolucao["despesa_pessoal_percentual"] / 100,
                None if evolucao.get("receitas_fundeb_percentual") is None else evolucao["receitas_fundeb_percentual"] / 100,
                None if evolucao.get("comprometimento_pontos_percentuais") is None else evolucao["comprometimento_pontos_percentuais"] / 100,
                " | ".join(item.get("avisos") or []),
                (item.get("limites") or {}).get("alerta", 0) / 100,
                (item.get("limites") or {}).get("prudencial", 0) / 100,
                (item.get("limites") or {}).get("maximo", 0) / 100,
            ]
        )
    _cabecalho(historico[1])
    historico.freeze_panes = "A2"
    historico.auto_filter.ref = f"A1:R{historico.max_row}"
    for linha in range(2, historico.max_row + 1):
        for coluna in (2, 3, 4, 8, 9):
            historico.cell(linha, coluna).number_format = MOEDA
        for coluna in (5, 6, 10, 11, 12, 13, 14, 16, 17, 18):
            historico.cell(linha, coluna).number_format = PERCENTUAL
        historico.cell(linha, 7).fill = PatternFill(
            "solid", fgColor=CORES_STATUS.get(historico.cell(linha, 7).value, CINZA)
        )
    larguras = [10, 20, 20, 22, 14, 16, 20, 20, 24, 16, 16, 16, 16, 23, 55, 16, 18, 16]
    for indice, largura in enumerate(larguras, start=1):
        historico.column_dimensions[historico.cell(1, indice).column_letter].width = largura

    grafico_valores = LineChart()
    grafico_valores.title = "Evolução dos valores financeiros"
    grafico_valores.y_axis.title = "Valor em R$"
    grafico_valores.x_axis.title = "Exercício"
    grafico_valores.add_data(
        Reference(historico, min_col=3, max_col=4, min_row=1, max_row=historico.max_row),
        titles_from_data=True,
    )
    grafico_valores.set_categories(Reference(historico, min_col=1, min_row=2, max_row=historico.max_row))
    grafico_valores.height = 8
    grafico_valores.width = 15
    resumo.add_chart(_sem_suavizar(grafico_valores), "F2")

    grafico_percentual = LineChart()
    grafico_percentual.title = "Comprometimento com pessoal"
    grafico_percentual.y_axis.title = "%"
    grafico_percentual.add_data(
        Reference(historico, min_col=6, min_row=1, max_row=historico.max_row),
        titles_from_data=True,
    )
    grafico_percentual.set_categories(Reference(historico, min_col=1, min_row=2, max_row=historico.max_row))
    grafico_percentual.height = 8
    grafico_percentual.width = 15
    resumo.add_chart(_sem_suavizar(grafico_percentual), "F18")

    caches = criar_painel(workbook, dados, resultados)

    erros = workbook.create_sheet("Erros e pendências")
    erros.append(["Ano", "Mensagem"])
    for erro in dados.get("erros") or []:
        erros.append([erro.get("ano"), erro.get("mensagem")])
    if erros.max_row == 1:
        erros.append(["—", "Nenhum exercício falhou na coleta."])
    _cabecalho(erros[1])
    erros.column_dimensions["A"].width = 14
    erros.column_dimensions["B"].width = 100
    for celula in erros["B"]:
        celula.alignment = Alignment(wrap_text=True, vertical="top")

    _configurar_impressao(resumo, "A1:T33", altura=2)
    _configurar_impressao(historico, f"A1:R{historico.max_row}")
    painel = workbook["Painel financeiro"]
    _configurar_impressao(
        painel,
        f"A1:{get_column_letter(3 + len(resultados))}48",
        altura=1,
    )
    painel.page_setup.paperSize = painel.PAPERSIZE_A3
    workbook.active = workbook.sheetnames.index("Painel financeiro")
    _configurar_impressao(erros, f"A1:B{erros.max_row}", paisagem=False)

    return injetar_cache(_finalizar(workbook), workbook.sheetnames, caches)


def gerar_excel(dados):
    if "resultados" in dados:
        return gerar_excel_historico(dados)
    return gerar_excel_anual(dados)
