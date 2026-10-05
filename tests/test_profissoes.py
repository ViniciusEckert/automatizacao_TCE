import json
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from app import create_app
from src.salarial.analise_service import analisar_tabela_salarial
from src.salarial.revisao_service import extrair_para_revisao, organizar_registros
from src.exportacao.excel_salarial import gerar_excel_salarial

ROOT=Path(__file__).resolve().parents[1]


class ProfissoesTestCase(unittest.TestCase):
    def setUp(self):
        self.dados=json.loads((ROOT/'data/processed/curitiba_administrativo_2026.json').read_text(encoding='utf-8'))

    def test_curitiba_preserva_184_vencimentos_reais(self):
        a=analisar_tabela_salarial(self.dados)
        self.assertEqual(a['resumo']['total_celulas'],184)
        self.assertEqual(a['tabela_atual'][0]['valores'][0],2180.44)
        self.assertEqual(a['tabela_atual'][-1]['valores'][-1],11490.21)
        self.assertEqual(a['resumo']['celulas_abaixo_base'],2)
        self.assertEqual(a['resumo']['celulas_abaixo_referencia'],184)
        self.assertEqual(a['referencia']['tipo'],'cenario')

    def test_csv_preserva_todas_as_classes(self):
        r=extrair_para_revisao('curitiba.csv',(ROOT/'data/modelos/curitiba_administrativo_2026.csv').read_bytes())
        self.assertEqual(len(r['candidatos']),1)
        d=r['candidatos'][0]
        self.assertEqual(d['classes'],self.dados['classes'])
        self.assertEqual([n['valores'] for n in d['niveis']],[n['valores'] for n in self.dados['niveis']])
        self.assertNotIn('resumo',r)

    def test_cargos_e_jornadas_nao_sao_misturados(self):
        rows=[dict(municipio='Teste',profissao=c,ano=2026,jornada=j,nivel='A',classe='1',vencimento=1000) for c,j in [('Analista',40),('Analista',30),('Técnico',40)]]
        self.assertEqual(len(organizar_registros(rows,'teste')),3)

    def test_duplicata_ou_lacuna_nao_e_descartada(self):
        row=dict(municipio='Teste',profissao='Analista',ano=2026,jornada=40,nivel='A',classe='1',vencimento=1000)
        with self.assertRaisesRegex(ValueError,'duplicado'):organizar_registros([row,row],'teste')
        with self.assertRaisesRegex(ValueError,'incompletas'):organizar_registros([row,{**row,'classe':'2'},{**row,'nivel':'B'}],'teste')

    def test_matriz_html_sem_inventar_metadados(self):
        html=b'<table><tr><th>Nivel</th><th>Descricao</th><th>1</th><th>2</th></tr><tr><td>A</td><td>Inicial</td><td>2000,00</td><td>2100,00</td></tr></table>'
        r=extrair_para_revisao('tabela.html',html)
        self.assertEqual(r['candidatos'][0]['niveis'][0]['valores'],[2000,2100])
        self.assertNotIn('ano',r['candidatos'][0])
        self.assertFalse(r['candidatos'][0]['jornada_confirmada_no_documento'])

    def test_outra_profissao_nao_herda_pspn(self):
        d=deepcopy(self.dados);d.pop('referencia');d['piso_nacional_40h']=5130.63
        with self.assertRaisesRegex(ValueError,'específica'):analisar_tabela_salarial(d)
        d=deepcopy(self.dados);d['referencia']['tipo']='pspn'
        with self.assertRaisesRegex(ValueError,'PSPN'):analisar_tabela_salarial(d)

    def test_44h_e_proporcionalidade_explicita(self):
        d=deepcopy(self.dados);d['jornada_semanal']=44;d['referencia']['jornada']=44
        self.assertEqual(analisar_tabela_salarial(d)['resumo']['vencimento_inicial_referencia'],2289.46)
        d['referencia']['jornada']=40
        with self.assertRaisesRegex(ValueError,'Jornadas diferentes'):analisar_tabela_salarial(d)
        d['referencia']['proporcional']=True
        self.assertEqual(analisar_tabela_salarial(d)['resumo']['vencimento_inicial_referencia'],2518.41)

    def test_ano_incompativel_e_nao_finitos(self):
        d=deepcopy(self.dados);d['referencia']['ano']=2025
        with self.assertRaisesRegex(ValueError,'ano'):analisar_tabela_salarial(d)
        for v in ['NaN','Infinity','-Infinity']:
            d=deepcopy(self.dados);d['referencia']['valor']=v
            with self.assertRaisesRegex(ValueError,'finito'):analisar_tabela_salarial(d)

    def test_exportacao_recalcula_e_mantem_cache(self):
        a=analisar_tabela_salarial(self.dados)
        a['resumo']['vencimento_inicial_referencia']=999999
        a['tabela_atual'][0]['valores'][0]=999999
        blob=gerar_excel_salarial(a)
        wb=load_workbook(BytesIO(blob),data_only=True)
        self.assertEqual(wb['Parâmetros']['B13'].value,2289.46)
        self.assertEqual(wb['Dados originais']['D4'].value,2180.44)
        self.assertEqual(wb['Dados originais']['AW7'].value,11490.21)
        formulas=load_workbook(BytesIO(blob),data_only=False)
        self.assertTrue(formulas['Comparação salarial']['D28'].value.startswith('='))
        self.assertIn('SUMPRODUCT',formulas['Comparação salarial']['A34'].value)

    def test_api_importacao_revisao_exportacao(self):
        with TemporaryDirectory() as temp:
            client=create_app(dossie_diretorio=temp).test_client()
            source=client.get('/api/salarios/amostras/curitiba-administrativo-2026')
            self.assertEqual(source.status_code,200)
            parsed=client.post('/api/salarios/extrair-revisao',data={'arquivo':(BytesIO((ROOT/'data/modelos/curitiba_administrativo_2026.csv').read_bytes()),'curitiba.csv')})
            self.assertEqual(parsed.status_code,200)
            self.assertTrue(parsed.json['revisao_obrigatoria'])
            self.assertNotIn('tabela_referencia',parsed.json)
            dados=source.json;dados['referencia']['valor']=2500
            result=client.post('/api/salarios/analisar',json=dados)
            self.assertEqual(result.status_code,200)
            self.assertEqual(result.json['resumo']['vencimento_inicial_referencia'],2500)
            excel=client.post('/api/salarios/exportar-excel',json=dados)
            self.assertEqual(excel.status_code,200)
            self.assertIn('administrativo',excel.headers['Content-Disposition'])
            self.assertEqual(load_workbook(BytesIO(excel.data),data_only=True)['Parâmetros']['B13'].value,2500)
            self.assertEqual(client.post('/api/salarios/relatorio-pdf',json=dados).status_code,200)


if __name__=='__main__':unittest.main()
