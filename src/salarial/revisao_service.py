"""Extração para revisão: nunca mistura profissões/jornadas nem compara sozinha."""
import csv
import json
import re
import unicodedata
from collections import OrderedDict
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from src.salarial.documento_service import importar_tabela_salarial, _numero_brasileiro
from src.salarial.analise_service import _decimal
from src.salarial.tabelas_coordenadas import extrair_pdf, extrair_html
from src.salarial.leitor_geral import ler_matrizes
from src.salarial.formatos_gerais import ler_pdf,ler_word,ler_xls,ler_imagem


def chave(v):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(v or '').strip().lower()) if not unicodedata.combining(c)).replace(' ','_')


class TabelasHTML(HTMLParser):
    def __init__(self):
        super().__init__();self.tables=[];self.table=None;self.row=None;self.cell=None
    def handle_starttag(self,tag,attrs):
        if tag=='table':self.table=[]
        if tag=='tr':self.row=[]
        if tag in ('td','th'):self.cell=[]
    def handle_data(self,data):
        if self.cell is not None:self.cell.append(data)
    def handle_endtag(self,tag):
        if tag in ('td','th') and self.cell is not None:
            if self.row is not None:self.row.append(' '.join(self.cell).strip())
            self.cell=None
        if tag=='tr' and self.row is not None:
            if self.table is not None:self.table.append(self.row)
            self.row=None
        if tag=='table' and self.table is not None:self.tables.append(self.table);self.table=None


def organizar_registros(linhas,nome):
    """Formato longo: uma linha por nível/classe, agrupado pela identidade completa."""
    grupos=OrderedDict()
    for numero,linha in enumerate(linhas,2):
        if not any(str(v or '').strip() for v in linha.values()):continue
        d={chave(k):v for k,v in linha.items()}
        aliases={'cargo':'profissao','valor':'vencimento','salario':'vencimento','referencia':'classe','jornada_semanal':'jornada'}
        for old,new in aliases.items():
            if old in d and new not in d:d[new]=d[old]
        required=['municipio','profissao','ano','jornada','nivel','classe','vencimento']
        faltas=[k for k in required if d.get(k) is None or str(d.get(k)).strip()=='']
        if faltas:raise ValueError(f"Linha {numero}: faltam {', '.join(faltas)}. Nenhuma linha foi descartada.")
        ident=tuple(str(d[k]).strip() for k in ['municipio','profissao','ano','jornada'])
        group=grupos.setdefault(ident,{'rows':[],'cells':set()})
        coord=(str(d['nivel']).strip().upper(),str(d['classe']).strip())
        if coord in group['cells']:raise ValueError(f"Linha {numero}: nível/classe duplicado em {ident[1]}.")
        group['cells'].add(coord)
        valor=float(_decimal(d['vencimento'],f'Vencimento da linha {numero}'))
        if valor<=0:raise ValueError(f"Linha {numero}: vencimento deve ser positivo.")
        group['rows'].append({**d,'nivel':coord[0],'classe':coord[1],'vencimento':valor})
    candidatos=[]
    for ident,g in grupos.items():
        cls=list(dict.fromkeys(r['classe'] for r in g['rows']))
        levels=list(dict.fromkeys(r['nivel'] for r in g['rows']))
        ns=[]
        for lev in levels:
            rs=[r for r in g['rows'] if r['nivel']==lev]; by={r['classe']:r['vencimento'] for r in rs}
            if set(by)!=set(cls):raise ValueError(f"{ident[1]}, nível {lev}: classes incompletas. Separe as estruturas ou complete a transcrição.")
            ns.append({'codigo':lev,'descricao':str(rs[0].get('descricao') or lev),'valores':[by[c] for c in cls]})
        base=ns[0]['valores'][0]
        for i,n in enumerate(ns):n['regra_vertical']={'tipo':'base' if i==0 else 'percentual','valor':round((n['valores'][0]/base-1)*100,4),'nivel_referencia':None if i==0 else ns[0]['codigo']}
        prog=[round((b/a-1)*100,4) for a,b in zip(ns[0]['valores'],ns[0]['valores'][1:])]
        candidatos.append({'municipio':ident[0],'profissao':ident[1],'ano':int(float(ident[2])),'jornada_semanal':float(_decimal(ident[3],'Jornada')),
            'jornada_confirmada_no_documento':True,'classes':cls,'niveis':ns,'progressoes_classes_percentuais':prog,
            'progressao_classes_origem':'inferida','arquivo_fonte':nome,'lei':str(g['rows'][0].get('lei') or 'Não informada'),
            'url_fonte_tabela':str(g['rows'][0].get('fonte_url') or ''),
            'avisos_importacao':['Progressões horizontais e verticais inferidas dos valores; revise a lei antes da comparação.']})
    return candidatos


def organizar_matriz(table,nome):
    """Matriz explícita Nível | Descrição (opcional) | classes, sem metadados presumidos."""
    for pos,row in enumerate(table):
        if not row or chave(row[0]) not in {'nivel','codigo'}:continue
        first=2 if len(row)>1 and chave(row[1]) in {'descricao','titulacao'} else 1
        headers=[str(v).strip() for v in row[first:] if v is not None and str(v).strip()]
        if not headers:continue
        records=[]
        for numero,values in enumerate(table[pos+1:],pos+2):
            if not any(v is not None and str(v).strip() for v in values):break
            if len(values)<first+len(headers):raise ValueError(f'Matriz, linha {numero}: faltam valores salariais.')
            code=str(values[0] or '').strip()
            if not code:raise ValueError(f'Matriz, linha {numero}: falta o nível.')
            for j,cl in enumerate(headers):
                records.append(dict(municipio='A informar',profissao='A informar',ano=2000,jornada=1,nivel=code,descricao=values[1] if first==2 else code,classe=cl,vencimento=values[first+j]))
        if records:
            d=organizar_registros(records,nome)[0]
            for key in ['municipio','profissao','ano','jornada_semanal']:d.pop(key,None)
            d['jornada_confirmada_no_documento']=False
            d['avisos_importacao'].append('Município, profissão, ano e jornada não constam no cabeçalho da matriz: informe-os na revisão.')
            return d
    return None


def extrair_para_revisao(nome,conteudo):
    if not conteudo or len(conteudo)>16*1024*1024:raise ValueError('Envie um arquivo não vazio de até 16 MB.')
    ext=Path(nome).suffix.lower(); candidates=[];text='';pendencias=[]
    if ext=='.json':
        obj=json.loads(conteudo.decode('utf-8-sig'))
        if isinstance(obj,dict) and 'niveis' in obj:candidates=[obj]
        else:candidates=organizar_registros(obj,nome)
    elif ext in {'.csv','.tsv','.xlsx','.xls','.html','.htm'}:
        tables=[]
        if ext in {'.csv','.tsv'}:
            try:text=conteudo.decode('utf-8-sig')
            except UnicodeDecodeError:text=conteudo.decode('cp1252')
            try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=';,\t')
            except csv.Error:dialect=csv.excel_tab if ext=='.tsv' else csv.excel
            tables=[list(csv.reader(StringIO(text),dialect))]
        elif ext=='.xls':
            tables=ler_xls(conteudo)
        elif ext=='.xlsx':
            wb=load_workbook(BytesIO(conteudo),read_only=True,data_only=True)
            for ws in wb:
                if ws.max_row>20000 or ws.max_column>220:raise ValueError('Planilha extensa: exporte somente a tabela selecionada.')
                tables.append(list(ws.values))
            wb.close()
        else:
            candidates,notes=extrair_html(nome,conteudo,organizar_registros);pendencias.extend(notes)
            parser=TabelasHTML();parser.feed(conteudo.decode('utf-8',errors='replace'));tables=[] if candidates else parser.tables
        rejected=[]
        for i,table in enumerate(tables,1):
            if not table:continue
            contexto='\n'.join(' '.join(str(c or '') for c in row) for row in table[:5])
            gerais,notas=ler_matrizes([table],nome,organizar_registros,contexto)
            if gerais:
                candidates.extend(gerais);continue
            if notas:
                pendencias.extend(notas);continue
            heads=[chave(c) for c in table[0]]
            if ('profissao' in heads or 'cargo' in heads) and ('vencimento' in heads or 'valor' in heads or 'salario' in heads):
                candidates.extend(organizar_registros([dict(zip(heads,r)) for r in table[1:]],nome))
            else:
                matrix=organizar_matriz(table,nome)
                if matrix:candidates.append(matrix)
                else:rejected.append(i)
        if rejected:pendencias.append(f'Tabelas/abas {rejected} não têm o cabeçalho do formato longo. Transcreva o recorte pelo modelo de importação; elas não foram incorporadas.')
    elif ext in {'.png','.jpg','.jpeg'}:
        candidates,text,notes=ler_imagem(nome,conteudo,organizar_registros);pendencias.extend(notes)
    elif ext=='.pdf':
        try:candidates,text,notes=extrair_pdf(nome,conteudo,organizar_registros)
        except ValueError:raise
        except Exception as erro:raise ValueError('PDF inválido, protegido ou sem estrutura legível; envie outro arquivo.') from erro
        pendencias.extend(notes)
        reconhecidas={int(m.group(1)) for d in candidates for m in [re.search(r'página (\d+)',d.get('arquivo_fonte',''))] if m}
        # Páginas rejeitadas pelo adaptador não são reinterpretadas como completas.
        reconhecidas.update(int(m.group(1)) for nota in notes for m in [re.search(r'página (\d+)',nota)] if m)
        gerais,texto_geral,notas=ler_pdf(nome,conteudo,organizar_registros,reconhecidas)
        candidates.extend(gerais);pendencias.extend(notas)
        if not text:text=texto_geral
        if not candidates and not pendencias:
            reader=PdfReader(BytesIO(conteudo))
            for i,pagina in enumerate(reader.pages,1):
                try:
                    parsed=importar_tabela_salarial('recorte.pdf',_pdf_pagina(reader,i-1))
                    parsed['arquivo_fonte']=f'{nome}, página {i}';candidates.append(parsed)
                except ValueError:pass
        pendencias.append('PDFs podem reunir carreiras ou continuações. Confirme o recorte completo; páginas não reconhecidas não foram incorporadas.')
    elif ext in {'.doc','.docx'}:
        candidates,text,notes=ler_word(nome,conteudo,organizar_registros);pendencias.extend(notes)
        if not candidates and not notes:
            try:candidates=[importar_tabela_salarial(nome,conteudo)]
            except ValueError as erro:pendencias.append(str(erro))
    else:raise ValueError('Use PDF, DOC/DOCX, XLS/XLSX, CSV/TSV, HTML, JSON ou PNG/JPG.')
    if not candidates:pendencias.append('Nenhuma tabela foi estruturada automaticamente. O texto disponível pode ser transcrito na grade ou no CSV; PDF imagem exige OCR.')
    return {'arquivo':nome,'candidatos':candidates,'texto_extraido':text[:250000],'texto_truncado':len(text)>250000,'pendencias':pendencias,'revisao_obrigatoria':True}


def _pdf_pagina(reader,index):
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_page(reader.pages[index]);out=BytesIO();writer.write(out);return out.getvalue()
