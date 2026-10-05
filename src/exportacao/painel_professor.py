"""Painel do professor: apresentação fiel, fontes e premissas explícitas."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MOEDA='R$ #,##0.00'
PERCENTUAL='0.00%'


def criar_painel(workbook, dados, resultados):
    ws=workbook.create_sheet('Painel financeiro')
    ps=workbook.create_sheet('Parâmetros do painel')
    ps.append(['Parâmetro','Valor','Premissa'])
    ps.append(['Média anual','exercicios','exercicios: acumulada/número de anos; intervalos: acumulada/distância entre anos; geometrica: taxa composta por intervalo.'])
    ps.append(['Exercícios analisados',len(resultados),'Não inclui anos sem resultado.'])
    ps.append(['Intervalos',resultados[-1]['ano']-resultados[0]['ano'],'Distância temporal entre primeiro e último exercício válido.'])
    ps.append(['Fonte RCL ajustada','RGF (relatório 20)','O modelo rotula RREO; a fonte efetivamente coletada é RGF, conferida com relatório 10.'])
    ps.append(['Interpretação MDE','não automatizada','A referência constitucional de 25% exige indicador e denominador confirmados.'])
    ps.append(['Primeiro semestre','não coletado','Não há valores mensais no coletor atual; nenhuma estimativa foi preenchida.'])
    dv=DataValidation(type='list',formula1='"exercicios,intervalos,geometrica"');ps.add_data_validation(dv);dv.add(ps['B2'])
    ps.append(['Ano','Alerta','Prudencial','Máximo'])
    ws.sheet_view.showGridLines=False
    end=1+len(resultados);last=get_column_letter(end);mean=get_column_letter(end+2)
    width=max(end+2,8);right=get_column_letter(width)
    for row,title in [(1,'PAINEL ORÇAMENTÁRIO/FISCAL'),(2,f"Município: {dados.get('municipio','')} - {dados.get('uf','PR')}")]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end)
        ws.cell(row,1,title);ws.cell(row,1).font=Font(name='Arial Black',size=14 if row==1 else 11,bold=True)
        ws.cell(row,1).fill=PatternFill('solid',fgColor='BFBFBF');ws.row_dimensions[row].height=26
    ws['A4']='Indicador'
    labels={6:'RREO - Receita Corrente Líquida Ajustada',9:'RGF - Despesa Total com Pessoal',12:'RGF - Despesa Total com Pessoal %',13:'RREO - Transferências FUNDEB',16:'Percentual de Aplicação em MDE\nsobre a Receita Líquida de Impostos\n(Referência constitucional: 25%; interpretação não automatizada)',17:'Classificação fiscal'}
    for row,label in labels.items():ws.cell(row,1,label)
    caches={};h=workbook['Histórico']
    for i,item in enumerate(resultados,2):
        col=get_column_letter(i);prev=get_column_letter(i-1);hr=i;pr=i+7
        ws.cell(4,i,item['ano']);ws.merge_cells(start_row=4,start_column=i,end_row=5,end_column=i)
        # Histórico already contains the official limits/defaults used by the analysis.
        vals=[h[f'{c}{hr}'].value for c in ('P','Q','R')]
        ps.append([item['ano']]+vals)
        for row,source in [(6,'C'),(9,'D'),(13,'H')]:
            value=h[f'{source}{hr}'].value
            formula=f'=IF(\'Histórico\'!{source}{hr}="","",\'Histórico\'!{source}{hr})'
            ws.cell(row,i,formula);caches[f'{col}{row}']='' if value is None else value
            ws.cell(row,i).number_format=MOEDA
            accum=row+1;annual=row+2
            if i==2:
                ws.cell(accum,i,f'=IF(OR({col}{row}="",{col}{row}=0),"",0)');caches[f'{col}{accum}']=0 if value not in (None,0) else ''
                ws.cell(annual,i,'Evolução')
            else:
                base=h[f'{source}2'].value;previous=h[f'{source}{hr-1}'].value
                ws.cell(accum,i,f'=IFERROR(IF(OR({col}{row}="",$B${row}=""),"",({col}{row}/$B${row})-1),"")')
                ws.cell(annual,i,f'=IFERROR(IF(OR({col}{row}="",{prev}{row}=""),"",({col}{row}/{prev}{row})-1),"")')
                caches[f'{col}{accum}']=value/base-1 if value is not None and base not in (None,0) else ''
                caches[f'{col}{annual}']=value/previous-1 if value is not None and previous not in (None,0) else ''
            ws.cell(accum,i).number_format=ws.cell(annual,i).number_format=PERCENTUAL
        ws.cell(12,i,f'=IFERROR(IF(OR({col}9="",{col}6=""),"",{col}9/{col}6),"")');caches[f'{col}12']=item['despesa_total_pessoal']/item['receita_corrente_liquida_ajustada']
        ws.cell(12,i).number_format=PERCENTUAL
        ws.cell(16,i,f'=IF(\'Histórico\'!J{hr}="","",\'Histórico\'!J{hr})');caches[f'{col}16']='' if h[f'J{hr}'].value is None else h[f'J{hr}'].value;ws.cell(16,i).number_format=PERCENTUAL
        a=f"'Parâmetros do painel'!B{pr}";p=f"'Parâmetros do painel'!C{pr}";m=f"'Parâmetros do painel'!D{pr}"
        ws.cell(17,i,f'=IF({col}12="","",IF({col}12<{a},"Normal",IF({col}12<{p},"Alerta",IF({col}12<{m},"Prudencial","Limite excedido"))))')
        v=caches[f'{col}12'];caches[f'{col}17']='Normal' if v<vals[0] else 'Alerta' if v<vals[1] else 'Prudencial' if v<vals[2] else 'Limite excedido'
        # Excel conditional formatting cannot directly refer to another sheet: local per-year limit cells.
        for row,val in zip((46,47,48),vals):
            ws.cell(row,i,f"='Parâmetros do painel'!{('B','C','D')[row-46]}{pr}");caches[f'{col}{row}']=val;ws.cell(row,i).number_format=PERCENTUAL
        for formula,color in [(f'AND(ISNUMBER({col}12),{col}12<{col}46)','EAF1DD'),(f'AND(ISNUMBER({col}12),{col}12>={col}46,{col}12<{col}47)','FFF2CC'),(f'AND(ISNUMBER({col}12),{col}12>={col}47,{col}12<{col}48)','FFE699'),(f'AND(ISNUMBER({col}12),{col}12>={col}48)','F4CCCC')]:
            ws.conditional_formatting.add(f'{col}12',FormulaRule(formula=[formula],fill=PatternFill('solid',fgColor=color)))
    for row in (7,10,14):
        ws.cell(row,end+1,'Média Anual')
        ws.cell(row,end+2,f'=IFERROR(IF(ISNUMBER({last}{row}),IF(\'Parâmetros do painel\'!B2="geometrica",(1+{last}{row})^(1/\'Parâmetros do painel\'!B4)-1,{last}{row}/IF(\'Parâmetros do painel\'!B2="intervalos",\'Parâmetros do painel\'!B4,\'Parâmetros do painel\'!B3)),""),"")')
        value=caches[f'{last}{row}'];caches[f'{mean}{row}']=value/len(resultados) if isinstance(value,(int,float)) else '';ws.cell(row,end+2).number_format=PERCENTUAL
    for row,title,color in [(18,'LIMITE MÁXIMO — art. 20 da LRF','F4CCCC'),(19,'LIMITE PRUDENCIAL — art. 22 da LRF','FFE699'),(20,'LIMITE DE ALERTA — art. 59 da LRF','BDD7EE')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2)
        ws.cell(row,1,title + f' — {vals[{18:2,19:1,20:0}[row]]:.2%} ({resultados[-1]["ano"]})');ws.cell(row,1).fill=PatternFill('solid',fgColor=color);ws.cell(row,1).alignment=Alignment(wrap_text=True)
    if end>=3:
        for row,text in [(19,'Fonte: SIM-AM, TCE-PR'),(20,f"Relatório emitido em: {dados.get('coletado_em') or 'data não informada'}")]:
            ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=end);ws.cell(row,3,text)
    else:ws['A21']='Fonte: SIM-AM, TCE-PR | '+str(dados.get('coletado_em') or 'data não informada')
    chart=BarChart();chart.type='col';chart.title='Indicadores financeiros';chart.y_axis.title='R$';chart.y_axis.numFmt='R$ #,##0,," mi"'
    for row in (6,9,13):chart.add_data(Reference(ws,min_col=1,max_col=end,min_row=row,max_row=row),titles_from_data=True,from_rows=True)
    chart.set_categories(Reference(ws,min_col=2,max_col=end,min_row=4,max_row=4));chart.width=25;chart.height=10;ws.add_chart(chart,'A23')
    notes=['RCL ajustada: o rótulo RREO reproduz o modelo; os valores vêm do RGF (relatório 20).','Média anual provisória: acumulada dividida pelos exercícios; altere em Parâmetros do painel.','MDE: sem classificação automática de conformidade. Primeiro semestre não coletado.']
    for item in resultados:
        notes.extend(item.get('avisos') or [])
    for error in dados.get('erros') or []:notes.append(f"{error.get('ano')}: {error.get('mensagem')}")
    ws.merge_cells(f'A42:{right}45');ws['A42']='\n'.join(dict.fromkeys(notes));ws['A42'].alignment=Alignment(wrap_text=True,vertical='top');ws['A42'].font=Font(name='Arial',size=10)
    for row,label in [(46,'Limite de alerta por exercício'),(47,'Limite prudencial por exercício'),(48,'Limite máximo por exercício')]:ws.cell(row,1,label)
    for row in ws.iter_rows(min_row=4,max_row=20,max_col=width):
        for c in row:
            if c.__class__.__name__=='MergedCell':continue
            c.font=Font(name='Arial',size=10,bold=c.column==1);c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='left' if c.column==1 else 'center')
            if c.row in (4,5):c.fill=PatternFill('solid',fgColor='000000');c.font=Font(name='Arial',color='FFFFFF',bold=True)
            elif c.row in (6,9,13):c.fill=PatternFill('solid',fgColor='D8D8D8')
            elif c.row==12:c.fill=PatternFill('solid',fgColor='EAF1DD')
    ws.column_dimensions['A'].width=48
    for i in range(2,width+1):ws.column_dimensions[get_column_letter(i)].width=24 if i<=end else 17
    for row in (6,9,12,13,16,18,19,20):ws.row_dimensions[row].height=40 if row!=16 else 60
    for row in range(42,46):ws.row_dimensions[row].height=19
    ws.freeze_panes='B6'
    ps.column_dimensions['A'].width=30;ps.column_dimensions['B'].width=25;ps.column_dimensions['C'].width=75;ps.column_dimensions['D'].width=18
    for row in ps:
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Arial',size=10)
    ps.row_dimensions[2].height=55;ps.row_dimensions[5].height=45;ps.row_dimensions[6].height=40
    ps.print_area=f'A1:D{ps.max_row}';ps.page_setup.orientation='landscape';ps.page_setup.fitToWidth=1;ps.page_setup.fitToHeight=1;ps.sheet_properties.pageSetUpPr.fitToPage=True
    return {'Painel financeiro':caches}
