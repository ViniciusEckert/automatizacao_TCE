from pathlib import Path
from openpyxl import load_workbook
import math,json
base=Path('work/v0.8.2/projeto_jornada_magisterio/docs/evidencias/v0.8.2');out=[];mismatch=[]
for name in ['financeiro','cafezal','curitiba']:
 f=load_workbook(base/(name+'.xlsx'),data_only=False);a=load_workbook(base/(name+'.xlsx'),data_only=True);b=load_workbook(Path('output/recalc082')/(name+'.xlsx'),data_only=True);count=0
 for ws in f:
  for row in ws:
   for c in row:
    if c.data_type!='f':continue
    count+=1;x=a[ws.title][c.coordinate].value;y=b[ws.title][c.coordinate].value
    if b[ws.title][c.coordinate].data_type=='e':mismatch.append([name,ws.title,c.coordinate,x,y,'error'])
    elif isinstance(x,(float,int)) and isinstance(y,(float,int)):
     if not math.isclose(x,y,rel_tol=1e-9,abs_tol=1e-7):mismatch.append([name,ws.title,c.coordinate,x,y])
    elif (x or '')!=(y or ''):mismatch.append([name,ws.title,c.coordinate,x,y])
 out.append({'arquivo':name,'formulas':count})
print(json.dumps({'resumo':out,'divergencias':mismatch[:30],'total':len(mismatch)},ensure_ascii=False,indent=2))
assert not mismatch
