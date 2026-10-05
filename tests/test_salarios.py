import unittest
from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from src.exportacao.excel_salarial import gerar_excel_salarial
from src.exportacao.relatorio_pdf import gerar_relatorio_salarial_pdf
from src.salarial.amostras import carregar_cafezal_2026
from src.salarial.analise_service import analisar_tabela_salarial
from src.salarial.documento_service import extrair_tabela_salarial_texto


TEXTO_ANEXO = """
Anexo A da LC 064/2026
Nível A - Magistério 2.565,31 2.616,61 2.668,94
Nível B - Licenciatura Plena 2.873,14 2.930,60 2.989,21
Nível C - Especialização 3.257,94 3.323,09 3.389,55
Percentual entre classes = 2,00%
Nível B = Nível A acrescido de 12,00%
Nível C = Nível A acrescido de 27,00%
"""


class SalariosTestCase(unittest.TestCase):
    def setUp(self):
        self.amostra = carregar_cafezal_2026()

    def test_truncamento_reproduz_anexo_de_cafezal(self):
        resultado = analisar_tabela_salarial(self.amostra)
        self.assertEqual(resultado["resumo"]["situacao"], "compativel")
        self.assertEqual(resultado["resumo"]["celulas_abaixo_referencia"], 0)
        self.assertEqual(resultado["resumo"]["divergencias_estrutura"], 0)
        self.assertEqual(resultado["tabela_referencia"][0]["valores"][0], 2565.31)
        self.assertEqual(resultado["tabela_referencia"][2]["valores"][6], 3668.94)

    def test_arredondamento_comercial_expoe_diferenca_de_centavos(self):
        dados = deepcopy(self.amostra)
        dados["modo_arredondamento"] = "arredondar"
        resultado = analisar_tabela_salarial(dados)
        self.assertEqual(resultado["tabela_referencia"][0]["valores"][0], 2565.32)
        self.assertGreater(resultado["resumo"]["celulas_abaixo_referencia"], 0)
        self.assertGreater(resultado["resumo"]["defasagem_inicial_percentual"], 0)

    def test_rejeita_modo_de_centavos_desconhecido(self):
        dados = deepcopy(self.amostra)
        dados["modo_arredondamento"] = "bancario"
        with self.assertRaisesRegex(ValueError, "truncar"):
            analisar_tabela_salarial(dados)

    def test_rejeita_niveis_com_quantidades_diferentes(self):
        dados = deepcopy(self.amostra)
        dados["niveis"][1]["valores"].pop()
        with self.assertRaisesRegex(ValueError, "mesma quantidade"):
            analisar_tabela_salarial(dados)

    def test_extrai_tabela_de_texto(self):
        resultado = extrair_tabela_salarial_texto(TEXTO_ANEXO, "anexo.docx")
        self.assertEqual(resultado["lei"], "Anexo A da LC 064/2026")
        self.assertEqual(resultado["progressao_classes_percentual"], 2)
        self.assertEqual(resultado["niveis"][1]["acrescimo_percentual"], 12)
        self.assertEqual(resultado["niveis"][2]["valores"][2], 3389.55)

    def test_excel_salarial_possui_formulas_e_fontes(self):
        analise = analisar_tabela_salarial(self.amostra)
        workbook = load_workbook(BytesIO(gerar_excel_salarial(analise)), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["Comparação salarial", "Parâmetros", "Dados originais", "Fontes e premissas"],
        )
        self.assertIn("ROUNDDOWN", workbook["Comparação salarial"]["E24"].value)
        self.assertEqual(workbook["Parâmetros"]["B9"].value, 5130.63)
        self.assertEqual(workbook["Parâmetros"]["B6"].value, 20)
        self.assertEqual(
            workbook["Comparação salarial"]["D42"].value,
            "=D16-D24",
        )
        self.assertIsNotNone(workbook["Parâmetros"]["B9"].comment)

    def test_relatorio_pdf_pode_ser_reaberto(self):
        analise = analisar_tabela_salarial(self.amostra)
        conteudo = gerar_relatorio_salarial_pdf(analise)
        self.assertTrue(conteudo.startswith(b"%PDF"))
        leitor = PdfReader(BytesIO(conteudo))
        self.assertGreaterEqual(len(leitor.pages), 4)
        texto = "\n".join(pagina.extract_text() or "" for pagina in leitor.pages)
        self.assertIn("Cafezal do Sul", texto)
        self.assertIn("Diferenças e premissas", texto)


if __name__ == "__main__":
    unittest.main()
