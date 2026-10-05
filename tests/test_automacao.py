import base64
from contextlib import ExitStack
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from pypdf import PdfWriter

from app import create_app
from src.coleta.tce_client import TCEError
from src.services.analise_service import AnaliseService
from test_service import ClienteFalso

ROOT = Path(__file__).resolve().parents[1]


class AutomacaoTestCase(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        ClienteFalso.chamadas = []
        self.app = create_app(AnaliseService(ClienteFalso, salvar_brutos=False), self.temp.name)
        self.app.config['TESTING'] = True
        self.client = self.app.test_client()
        self.patches = self.enterContext(ExitStack())
        self.patches.enter_context(patch('app._municipios_tce', return_value=(('1467', 'CRUZEIRO DO SUL'),)))
        self.patches.enter_context(patch('app._entidades_tce', return_value=(('9777', 'CÂMARA MUNICIPAL DE CRUZEIRO DO SUL'), ('12266', 'MUNICÍPIO DE CRUZEIRO DO SUL'))))
        self.relatorios = self.patches.enter_context(patch('app._relatorios_tce', return_value=(('10', 'RCL'), ('20', 'Pessoal'), ('15', 'MDE'))))
        self.anos = self.patches.enter_context(patch('app._anos_tce', return_value=(('2023', '2023'), ('2025', '2025'), ('2099', '2099'))))

    def workbook(self, resposta):
        self.assertEqual(resposta.status_code, 200, resposta.get_json())
        arquivo = resposta.json['arquivo']
        self.assertTrue(arquivo['nome'].endswith('.xlsx'))
        return load_workbook(BytesIO(base64.b64decode(arquivo['conteudo_base64'])), data_only=True)

    def preparar(self, nome='tabela.csv', conteudo=None):
        if conteudo is None:
            conteudo = (ROOT / 'data/modelos/curitiba_administrativo_2026.csv').read_bytes()
        resposta = self.client.post('/api/automacao/preparar-tabela', data={'arquivo': (BytesIO(conteudo), nome)})
        self.assertEqual(resposta.status_code, 200, resposta.json)
        return resposta.json

    def test_financeiro_somente_municipio_e_periodo_entrega_excel(self):
        resposta = self.client.post('/api/automacao/financeiro', json={'municipio_id': '1467', 'periodo': 'ultimo', 'entidade_id': '9777', 'municipio_nome': 'Nome não oficial'})
        wb = self.workbook(resposta)
        self.relatorios.assert_called_once_with('1467', '12266')
        self.anos.assert_called_once_with('1467', '12266')
        self.assertEqual(resposta.json['titulo'], 'CRUZEIRO DO SUL')
        self.assertEqual(resposta.json['periodo'], '2025')
        self.assertEqual({(r, a) for r, a, _ in ClienteFalso.chamadas}, {('10', 2025), ('20', 2025), ('15', 2025)})
        self.assertGreater(len(wb.sheetnames), 1)

    def test_historico_nao_consulta_ano_ausente_nem_inventa_valor(self):
        resposta = self.client.post('/api/automacao/financeiro', json={'municipio_id': '1467'})
        self.workbook(resposta)
        self.assertEqual(resposta.json['periodo'], '2023-2025')
        self.assertEqual({a for _, a, _ in ClienteFalso.chamadas}, {2023, 2025})
        self.assertTrue(any('2024' in aviso and 'publicação' in aviso for aviso in resposta.json['avisos']))

    def test_indisponibilidade_nao_tenta_a_camara_nem_baixa_falso_excel(self):
        self.relatorios.return_value = (('20', 'Pessoal'),)
        with self.assertLogs(self.app.logger, level='WARNING'):
            resposta = self.client.post('/api/automacao/financeiro', json={'municipio_id': '1467'})
        self.assertEqual(resposta.status_code, 422)
        self.assertNotIn('arquivo', resposta.json)
        self.assertEqual(ClienteFalso.chamadas, [])
        self.relatorios.assert_called_once_with('1467', '12266')

    def test_falha_da_fonte_tem_mensagem_simples_e_preserva_diagnostico(self):
        self.anos.side_effect = TCEError('Falha de teste no portal')
        with self.assertLogs(self.app.logger, level='WARNING'):
            resposta = self.client.post('/api/automacao/financeiro', json={'municipio_id': '1467'})
        self.assertEqual(resposta.status_code, 502)
        self.assertIn('Tente novamente', resposta.json['erro'])
        self.assertIn('Falha de teste no portal', resposta.json['detalhe'])
        self.assertNotIn('arquivo', resposta.json)

    def test_periodo_ou_municipio_invalido_nao_dispara_coleta(self):
        for pedido in ({'municipio_id': 'inventado'}, {'municipio_id': '1467', 'periodo': 'intervalo', 'ano_inicial': 2025, 'ano_final': 2024}):
            resposta = self.client.post('/api/automacao/financeiro', json=pedido)
            self.assertEqual(resposta.status_code, 400)
        self.assertEqual(ClienteFalso.chamadas, [])

    def test_exemplo_offline_identificado_e_sem_chamadas_ao_tce(self):
        resposta = self.client.get('/api/automacao/exemplo-financeiro')
        self.workbook(resposta)
        self.assertTrue(resposta.json['exemplo'])
        self.assertEqual(resposta.json['periodo'], '2019-2025')
        self.relatorios.assert_not_called()

    def test_duas_tabelas_preparadas_geram_modelo_completo_em_um_pedido(self):
        for perfil, linhas, colunas, primeiro, ultimo in [('cafezal-2026', 3, 12, 2565.31, 4050.77), ('curitiba-administrativo-2026', 4, 46, 2180.44, 11490.21)]:
            with self.subTest(perfil=perfil):
                resposta = self.client.post('/api/automacao/salarios', json={'tabela_id': perfil})
                wb = self.workbook(resposta)
                self.assertEqual(wb.sheetnames, ['Comparação salarial', 'Parâmetros', 'Dados originais', 'Fontes e premissas'])
                ws = wb['Dados originais']
                self.assertEqual(ws.cell(4, 4).value, primeiro)
                self.assertEqual(ws.cell(3 + linhas, 3 + colunas).value, ultimo)
                valores = [c.value for row in ws.iter_rows(min_row=4, max_row=3 + linhas, min_col=4, max_col=3 + colunas) for c in row]
                self.assertEqual(len(valores), linhas * colunas)
                self.assertTrue(all(isinstance(v, (int, float)) for v in valores))

    def test_documento_preparado_uma_vez_reutilizado_apos_reabrir(self):
        preparado = self.preparar()
        self.assertFalse(preparado['requer_conferencia_recorte'])
        pedido = {'preparacao_id': preparado['preparacao_id'], 'comparacao': {'tipo': 'percentual', 'valor': '5'}}
        resposta = self.client.post('/api/automacao/salarios', json=pedido)
        wb = self.workbook(resposta)
        self.assertEqual(wb['Parâmetros']['B13'].value, 2289.46)
        perfil = resposta.json['tabela']['id']
        nova_app = create_app(dossie_diretorio=self.temp.name)
        novo_client = nova_app.test_client()
        self.assertIn(perfil, [t['id'] for t in novo_client.get('/api/automacao/tabelas').json])
        repetida = novo_client.post('/api/automacao/salarios', json={'tabela_id': perfil})
        outro = self.workbook(repetida)
        self.assertEqual(outro['Dados originais']['AW7'].value, 11490.21)
        self.assertEqual(outro['Parâmetros']['B13'].value, 2289.46)
        dados = nova_app.config['AUTOMACAO_SALARIAL'].carregar(perfil)
        self.assertEqual(len(dados['documento_sha256']), 64)
        self.assertEqual(dados['referencia']['tipo'], 'cenario')
        self.assertTrue(dados['jornada_confirmada_no_documento'])

    def test_matriz_pede_dados_ausentes_e_nao_os_apresenta_como_extraidos(self):
        preparado = self.preparar('matriz.csv', b'Nivel;1;2\nA;2000;2100\n')
        d = preparado['candidatos'][0]
        self.assertNotIn('municipio', d)
        pedido = {'preparacao_id': preparado['preparacao_id'], 'comparacao': {'tipo': 'valor', 'valor': 2400}}
        resposta = self.client.post('/api/automacao/salarios', json=pedido)
        self.assertEqual(resposta.status_code, 400)
        self.assertIn('município', resposta.json['erro'])
        pedido['complementos'] = {'municipio': 'Cidade Teste', 'profissao': 'Técnico', 'ano': 2026, 'jornada_semanal': 40}
        resposta = self.client.post('/api/automacao/salarios', json=pedido)
        self.workbook(resposta)
        self.assertIn('premissa', resposta.json['tabela']['descricao'])
        self.assertTrue(any('carga horária' in a for a in resposta.json['avisos']))

    def test_piso_nao_aplicado_a_outra_profissao_e_ano_nao_presumido(self):
        preparado = self.preparar()
        resposta = self.client.post('/api/automacao/salarios', json={'preparacao_id': preparado['preparacao_id'], 'comparacao': {'tipo': 'pspn'}})
        self.assertEqual(resposta.status_code, 400)
        self.assertIn('PSPN', resposta.json['erro'])
        matriz = self.preparar('matriz.csv', b'Nivel;1\nA;2000\n')
        for ano in (2026.5, True, 'abc'):
            resposta = self.client.post('/api/automacao/salarios', json={'preparacao_id': matriz['preparacao_id'], 'complementos': {'municipio': 'Teste', 'profissao': 'Técnico', 'ano': ano, 'jornada_semanal': 40}, 'comparacao': {'tipo': 'valor', 'valor': 2000}})
            self.assertEqual(resposta.status_code, 400)

    def test_documento_multipagina_nao_exporta_recorte_sem_confirmar(self):
        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        writer.add_blank_page(100, 100)
        arquivo = BytesIO()
        writer.write(arquivo)
        candidato = json.loads((ROOT / 'data/processed/cafezal_do_sul_2026.json').read_text(encoding='utf-8'))
        with patch('src.salarial.automacao_service.extrair_para_revisao', return_value={'candidatos': [deepcopy(candidato)], 'pendencias': ['Confira a continuação.']}):
            preparado = self.preparar('duas-paginas.pdf', arquivo.getvalue())
        self.assertTrue(preparado['requer_conferencia_recorte'])
        pedido = {'preparacao_id': preparado['preparacao_id']}
        resposta = self.client.post('/api/automacao/salarios', json=pedido)
        self.assertEqual(resposta.status_code, 400)
        self.assertIn('completa', resposta.json['erro'])
        self.workbook(self.client.post('/api/automacao/salarios', json={**pedido, 'confirmar_recorte': True}))

    def test_json_malformado_nao_chega_ao_preview_e_nao_salva_perfil(self):
        resposta = self.client.post('/api/automacao/preparar-tabela', data={'arquivo': (BytesIO(b'{"niveis":[{"codigo":"A","valores":null}]}'), 'invalido.json')})
        self.assertEqual(resposta.status_code, 400)
        self.assertEqual(len(self.client.get('/api/automacao/tabelas').json), 2)

    def test_abertura_automatica_usa_porta_reservada_sem_debug(self):
        import iniciar
        servidor = Mock(server_port=51234)
        servidor.serve_forever.side_effect = KeyboardInterrupt
        with patch('iniciar.make_server', return_value=servidor) as factory, patch('iniciar.Timer') as timer, patch('iniciar.webbrowser.open') as browser, patch('builtins.print'):
            iniciar.iniciar()
        self.assertEqual(factory.call_args.args[:2], ('127.0.0.1', 0))
        self.assertTrue(factory.call_args.kwargs['threaded'])
        timer.assert_called_once_with(0.7, browser, args=('http://127.0.0.1:51234/',))
        timer.return_value.start.assert_called_once()
        servidor.server_close.assert_called_once()


if __name__ == '__main__':
    unittest.main()
