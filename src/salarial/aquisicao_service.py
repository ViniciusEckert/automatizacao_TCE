"""Aquisição auditável dos documentos brutos do módulo salarial."""

from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from openpyxl import load_workbook

from src.salarial.pdf_service import inspecionar_pdf


LIMITE_ARQUIVO_BYTES = 15 * 1024 * 1024
EXTENSOES_ACEITAS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".html", ".htm"}
CATEGORIAS = {
    "tabela_vigente": "Tabela salarial vigente",
    "plano_carreira": "Plano de carreira/estatuto",
    "ato_reajuste": "Lei, decreto ou portaria de reajuste",
}


class _ExtratorHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.partes = []

    def handle_data(self, data):
        texto = data.strip()
        if texto:
            self.partes.append(texto)


def _validar_url(url):
    valor = str(url or "").strip()
    partes = urlparse(valor)
    if partes.scheme not in {"http", "https"} or not partes.netloc:
        raise ValueError("Informe a URL pública onde o documento foi encontrado.")
    return valor


def _validar_arquivo(nome_arquivo, conteudo):
    nome = str(nome_arquivo or "").strip()
    extensao = Path(nome).suffix.lower()
    if not nome or extensao not in EXTENSOES_ACEITAS:
        formatos = ", ".join(sorted(EXTENSOES_ACEITAS))
        raise ValueError(f"Formato não aceito no dossiê. Use: {formatos}.")
    if not conteudo:
        raise ValueError(f"O arquivo {nome} está vazio.")
    if len(conteudo) > LIMITE_ARQUIVO_BYTES:
        raise ValueError(f"O arquivo {nome} excede o limite de 15 MB.")
    return nome, extensao


def _texto_docx(conteudo):
    try:
        documento = Document(BytesIO(conteudo))
    except Exception as erro:
        raise ValueError("O DOCX não pôde ser aberto.") from erro
    textos = [paragrafo.text.strip() for paragrafo in documento.paragraphs if paragrafo.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            textos.extend(celula.text.strip() for celula in linha.cells if celula.text.strip())
    return "\n".join(textos)


def _inspecionar_xlsx(conteudo):
    try:
        workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=False)
    except Exception as erro:
        raise ValueError("O XLSX não pôde ser aberto.") from erro
    celulas_preenchidas = 0
    for planilha in workbook.worksheets:
        for linha in planilha.iter_rows():
            celulas_preenchidas += sum(celula.value is not None for celula in linha)
    return len(workbook.sheetnames), celulas_preenchidas


def inspecionar_documento_bruto(nome_arquivo, conteudo, categoria):
    """Identifica se o documento pode seguir para triagem, sem extrair salários."""

    if categoria not in CATEGORIAS:
        raise ValueError("Categoria documental inválida.")
    nome, extensao = _validar_arquivo(nome_arquivo, conteudo)
    resultado = {
        "categoria": categoria,
        "categoria_descricao": CATEGORIAS[categoria],
        "nome_original": nome,
        "extensao": extensao,
        "tamanho_bytes": len(conteudo),
        "sha256": sha256(conteudo).hexdigest(),
        "texto_extraivel": None,
        "quantidade_caracteres": None,
        "requer_ocr": False,
        "requer_conversao": False,
        "status_inspecao": "registrado_para_triagem",
    }

    if extensao == ".pdf":
        pdf = inspecionar_pdf(nome, conteudo)
        resultado.update(
            texto_extraivel=pdf["texto_extraivel"],
            quantidade_caracteres=pdf["quantidade_caracteres"],
            requer_ocr=pdf["provavelmente_escaneado"],
            quantidade_paginas=pdf["quantidade_paginas"],
        )
    elif extensao == ".docx":
        texto = _texto_docx(conteudo)
        resultado.update(texto_extraivel=bool(texto), quantidade_caracteres=len(texto))
    elif extensao == ".xlsx":
        abas, celulas = _inspecionar_xlsx(conteudo)
        resultado.update(quantidade_abas=abas, celulas_preenchidas=celulas)
    elif extensao in {".csv", ".html", ".htm"}:
        texto = conteudo.decode("utf-8", errors="replace")
        if extensao in {".html", ".htm"}:
            extrator = _ExtratorHTML()
            extrator.feed(texto)
            texto = "\n".join(extrator.partes)
        resultado.update(texto_extraivel=bool(texto.strip()), quantidade_caracteres=len(texto.strip()))
    else:
        resultado["requer_conversao"] = True
        resultado["status_inspecao"] = "registrado_aguardando_conversao"

    return resultado


def validar_metadados_dossie(municipio, uf, ano, fonte_url, tipo_fonte, observacoes=""):
    municipio = str(municipio or "").strip()
    if len(municipio) < 2 or len(municipio) > 120:
        raise ValueError("Informe o município do dossiê.")
    uf = str(uf or "").strip().upper()
    if len(uf) != 2 or not uf.isalpha():
        raise ValueError("Informe a UF com duas letras.")
    try:
        ano = int(ano)
    except (TypeError, ValueError) as erro:
        raise ValueError("Informe o ano de referência do documento.") from erro
    if ano < 2000 or ano > 2100:
        raise ValueError("O ano de referência está fora do intervalo aceito.")
    tipo_fonte = str(tipo_fonte or "").strip()
    if tipo_fonte not in {"portal_prefeitura", "diario_oficial", "outro_oficial"}:
        raise ValueError("Informe o tipo de fonte oficial.")
    observacoes = str(observacoes or "").strip()
    if len(observacoes) > 2000:
        raise ValueError("As observações devem ter no máximo 2.000 caracteres.")
    return {
        "municipio": municipio,
        "uf": uf,
        "ano_referencia": ano,
        "fonte_url": _validar_url(fonte_url),
        "tipo_fonte": tipo_fonte,
        "observacoes": observacoes,
    }


def registrar_dossie_salarial(store, metadados, arquivos):
    """Valida todos os itens antes de criar qualquer arquivo no disco."""

    dados = validar_metadados_dossie(**metadados)
    recebidos = []
    for item in arquivos:
        conteudo = item["conteudo"]
        inspecao = inspecionar_documento_bruto(
            item["nome_original"], conteudo, item["categoria"]
        )
        recebidos.append({**inspecao, "conteudo": conteudo})
    categorias = {item["categoria"] for item in recebidos}
    if "tabela_vigente" not in categorias:
        raise ValueError("A tabela salarial vigente é obrigatória para criar o dossiê.")
    return store.salvar(dados, recebidos)
