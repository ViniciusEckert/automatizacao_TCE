"""Importação de tabelas salariais em PDF, DOCX e DOC legado."""

from io import BytesIO
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from docx import Document
from pypdf import PdfReader


PADRAO_MOEDA = re.compile(
    r"(?<![\d.])(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}(?!\d)(?!\s*%)"
)
CODIGO_NIVEL = r"[A-Z]+|\d+"
PADRAO_NIVEL = re.compile(
    rf"\bN[ií]vel\s+({CODIGO_NIVEL})\s*(?:[-–—:]\s*)",
    re.IGNORECASE,
)
PADRAO_PERCENTUAL = re.compile(r"([\d]+(?:[.,][\d]+)?)\s*%")
PADRAO_GATILHO_PROGRESSAO = re.compile(
    r"percentual\s+entre\s+(?:as\s+)?classes|"
    r"interst[ií]ci|"
    r"progress[aã]o(?:\s+horizontal)?[^\n]*classes|"
    r"classes[^\n]*progress[aã]o",
    flags=re.IGNORECASE,
)


PADRAO_VIGENCIA = re.compile(
    r"revog\w*|nova\s+reda[cç][aã]o|passa(?:m)?\s+a\s+vigorar|fica(?:m)?\s+alterad\w*|"
    r"(?:foi|foram|ser[aá]|ser[aã]o)\s+alterad\w*",
    re.IGNORECASE,
)


def _avisos_vigencia(texto):
    """Sinaliza termos de revogação/alteração; o sistema nunca decide qual redação está vigente."""
    trechos = []
    for linha in str(texto).splitlines():
        achado = PADRAO_VIGENCIA.search(linha)
        if achado:
            inicio = max(0, achado.start() - 40)
            trecho = " ".join(linha[inicio:achado.end() + 60].split())
            if trecho not in trechos:
                trechos.append(trecho)
    if not trechos:
        return []
    exemplos = "; ".join(f"«{t}»" for t in trechos[:3])
    return [
        "O documento contém termos de revogação ou alteração de regras "
        f"({len(trechos)} trecho(s), por exemplo: {exemplos}). O sistema não decide qual redação está vigente: "
        "confirme na lei que a tabela e as regras importadas são as vigentes."
    ]


def _numero_brasileiro(texto):
    return float(texto.replace(".", "").replace(",", "."))


def _texto_pdf(conteudo):
    try:
        leitor = PdfReader(BytesIO(conteudo))
    except Exception as erro:
        raise ValueError("O PDF não pôde ser aberto.") from erro
    paginas = [(pagina.extract_text() or "").strip() for pagina in leitor.pages]
    texto = "\n".join(pagina for pagina in paginas if pagina)
    if not texto.strip():
        raise ValueError("O PDF não possui texto extraível; execute OCR antes da importação.")
    return texto


def _texto_docx(conteudo):
    try:
        documento = Document(BytesIO(conteudo))
    except Exception as erro:
        raise ValueError("O DOCX não pôde ser aberto.") from erro
    linhas = [paragrafo.text for paragrafo in documento.paragraphs if paragrafo.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            linhas.append(" ".join(celula.text.strip() for celula in linha.cells))
    return "\n".join(linhas)


def _decodificar_saida(conteudo):
    for codificacao in ("utf-8", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("utf-8", errors="replace")


def _texto_doc_legado(conteudo):
    """Converte `.doc` com LibreOffice ou antiword, quando disponível."""

    with tempfile.TemporaryDirectory(prefix="jornada-doc-") as diretorio:
        pasta = Path(diretorio)
        origem = pasta / "tabela.doc"
        origem.write_bytes(conteudo)

        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            perfil = pasta / "perfil-libreoffice"
            perfil.mkdir()
            processo = subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation={perfil.as_uri()}",
                    "--headless",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    str(pasta),
                    str(origem),
                ],
                capture_output=True,
                timeout=45,
                check=False,
            )
            convertido = pasta / "tabela.txt"
            if processo.returncode == 0 and convertido.exists():
                return _decodificar_saida(convertido.read_bytes())

        antiword = shutil.which("antiword")
        if antiword:
            processo = subprocess.run(
                [antiword, str(origem)], capture_output=True, timeout=45, check=False
            )
            if processo.returncode == 0 and processo.stdout:
                return _decodificar_saida(processo.stdout)

    raise ValueError(
        "Para importar arquivos .doc legados, instale o LibreOffice ou converta o arquivo para DOCX/PDF."
    )


def extrair_texto_documento(nome_arquivo, conteudo):
    """Extrai texto preservando linhas suficientes para reconhecer a tabela."""

    if not conteudo:
        raise ValueError("O arquivo enviado está vazio.")
    if len(conteudo) > 16 * 1024 * 1024:
        raise ValueError("O arquivo excede o limite de 16 MB.")
    extensao = Path(nome_arquivo or "").suffix.lower()
    if extensao == ".pdf":
        return _texto_pdf(conteudo)
    if extensao == ".docx":
        return _texto_docx(conteudo)
    if extensao == ".doc":
        return _texto_doc_legado(conteudo)
    raise ValueError("Formato não aceito. Envie PDF, DOCX ou DOC.")


def _percentuais_progressao(texto, quantidade_classes):
    """Localiza uma regra horizontal uniforme ou uma lista por transição."""

    quantidade_transicoes = max(quantidade_classes - 1, 0)
    if quantidade_transicoes == 0:
        return [0.0], [], "declarada"

    declaracoes = []
    for numero_linha, linha in enumerate(texto.splitlines(), start=1):
        gatilho = PADRAO_GATILHO_PROGRESSAO.search(linha)
        if not gatilho:
            continue

        # Lê somente a cláusula iniciada pelo gatilho. Uma observação entre
        # parênteses após a regra não deve acrescentar percentuais à estrutura.
        trecho = linha[gatilho.start() :].split("(", maxsplit=1)[0]
        valores = [
            _numero_brasileiro(item.group(1))
            for item in PADRAO_PERCENTUAL.finditer(trecho)
        ]
        if not valores:
            continue

        if len(valores) == 1 or len(set(valores)) == 1:
            valor = valores[0]
            regra = ([valor], [valor] * quantidade_transicoes)
        elif len(valores) == quantidade_transicoes:
            regra = (valores, valores)
        else:
            raise ValueError(
                f"A progressão entre classes ficou ambígua na linha {numero_linha}: "
                "informe um percentual único ou "
                f"{quantidade_transicoes} percentuais, um para cada transição."
            )
        declaracoes.append((numero_linha, valores, regra))

    if declaracoes:
        regras_distintas = {
            (tuple(regra[0]), tuple(regra[1])) for _linha, _valores, regra in declaracoes
        }
        if len(regras_distintas) > 1:
            detalhes = "; ".join(
                f"linha {numero}: {', '.join(f'{valor:g}%' for valor in valores)}"
                for numero, valores, _regra in declaracoes
            )
            raise ValueError(
                "Foram encontradas declarações conflitantes de progressão entre classes "
                f"({detalhes}). Confirme qual regra está vigente na lei."
            )
        declarada, expandida = declaracoes[0][2]
        return declarada, expandida, "declarada"

    return None, None, None


def _inferir_progressoes(valores):
    percentuais = []
    for anterior, atual in zip(valores, valores[1:]):
        percentual = (atual / anterior - 1) * 100
        percentuais.append(round(percentual, 4))
    return percentuais


def _regras_verticais(texto, codigos):
    """Extrai regras entre níveis, aceitando percentual ou valor fixo em reais."""

    token = rf"({CODIGO_NIVEL})"
    referencia = rf"(?:N[ií]vel\s+({CODIGO_NIVEL})|(N[ií]vel\s+anterior))"
    inicio = rf"N[ií]vel\s+{token}\s*(?:=|corresponde\s+ao?)\s*{referencia}"
    percentual = re.compile(
        inicio
        + r"\s*(?:acrescid[oa]\s+de|mais|\+)\s*([\d]+(?:[.,][\d]+)?)\s*%",
        flags=re.IGNORECASE,
    )
    valor_fixo = re.compile(
        inicio
        + r"\s*(?:acrescid[oa]\s+de|mais|\+)\s*R\$\s*"
        + r"((?:\d{1,3}(?:\.\d{3})+|\d+),\d{2})",
        flags=re.IGNORECASE,
    )

    regras = {}
    posicoes = {codigo: indice for indice, codigo in enumerate(codigos)}
    for tipo, padrao in (("percentual", percentual), ("valor_fixo", valor_fixo)):
        for encontrado in padrao.finditer(texto):
            alvo = encontrado.group(1).upper()
            referencia_explicita = encontrado.group(2)
            usa_anterior = encontrado.group(3)
            valor = _numero_brasileiro(encontrado.group(4))
            if alvo not in posicoes or posicoes[alvo] == 0:
                continue
            if usa_anterior:
                codigo_referencia = codigos[posicoes[alvo] - 1]
            else:
                codigo_referencia = str(referencia_explicita).upper()
            if codigo_referencia not in posicoes or posicoes[codigo_referencia] >= posicoes[alvo]:
                raise ValueError(
                    f"A regra do nível {alvo} referencia um nível inexistente ou posterior."
                )
            nova_regra = {
                "tipo": tipo,
                "valor": valor,
                "nivel_referencia": codigo_referencia,
            }
            regra_anterior = regras.get(alvo)
            if regra_anterior is not None and regra_anterior != nova_regra:
                raise ValueError(
                    f"Foram encontradas regras verticais conflitantes para o nível {alvo}. "
                    "Confirme qual regra está vigente na lei."
                )
            regras[alvo] = nova_regra
    return regras


def extrair_tabela_salarial_texto(texto, nome_arquivo="documento"):
    """Reconhece níveis, classes e percentuais de um anexo textual."""

    niveis = []
    linhas_nao_reconhecidas = []
    codigos = set()
    for numero_linha, linha in enumerate(texto.splitlines(), start=1):
        valores_na_linha = list(PADRAO_MOEDA.finditer(linha))
        nivel_encontrado = PADRAO_NIVEL.search(linha)
        menciona_nivel = re.search(r"\bN[ií]vel\b", linha, flags=re.IGNORECASE)
        if menciona_nivel and len(valores_na_linha) >= 2 and not nivel_encontrado:
            linhas_nao_reconhecidas.append((numero_linha, re.sub(r"\s+", " ", linha).strip()))
            continue
        valores_encontrados = (
            list(PADRAO_MOEDA.finditer(linha, nivel_encontrado.end()))
            if nivel_encontrado
            else []
        )
        if not nivel_encontrado or len(valores_encontrados) < 2:
            continue
        codigo = nivel_encontrado.group(1).upper()
        if codigo in codigos:
            raise ValueError(f"O código de nível {codigo} aparece mais de uma vez na tabela.")
        codigos.add(codigo)
        inicio_valores = valores_encontrados[0].start()
        descricao = linha[nivel_encontrado.end() : inicio_valores]
        descricao = re.sub(r"\s+", " ", descricao).strip(" -") or f"Nível {codigo}"
        valores = [_numero_brasileiro(item.group(0)) for item in valores_encontrados]
        niveis.append(
            {
                "codigo": codigo,
                "descricao": descricao,
                "acrescimo_percentual": 0.0,
                "acrescimo_valor_fixo": None,
                "regra_vertical": {
                    "tipo": "base" if not niveis else "percentual",
                    "valor": 0.0,
                    "nivel_referencia": None if not niveis else niveis[0]["codigo"],
                },
                "valores": valores,
            }
        )

    if linhas_nao_reconhecidas:
        exemplos = "; ".join(
            f"linha {numero}: {linha[:100]}" for numero, linha in linhas_nao_reconhecidas[:3]
        )
        raise ValueError(
            "Foram encontradas linhas salariais com 'Nível' que não puderam ser interpretadas "
            f"com segurança ({exemplos}). Nenhuma linha foi descartada silenciosamente."
        )
    if not niveis:
        raise ValueError("Nenhuma linha salarial com níveis e valores foi reconhecida.")
    quantidade_classes = len(niveis[0]["valores"])
    if any(len(nivel["valores"]) != quantidade_classes for nivel in niveis):
        raise ValueError("As linhas salariais reconhecidas possuem quantidades diferentes de classes.")

    avisos = []
    _progressao_encontrada, progressoes, origem_progressao = _percentuais_progressao(
        texto, quantidade_classes
    )
    if progressoes is None:
        progressoes = _inferir_progressoes(niveis[0]["valores"])
        origem_progressao = "inferida"
        avisos.append(
            "As progressões entre classes foram inferidas matematicamente da primeira linha; "
            "confirme os percentuais na lei."
        )

    uniforme = not progressoes or len(set(progressoes)) == 1
    progressao = progressoes[0] if uniforme and progressoes else 0.0

    base = niveis[0]["valores"][0]
    regras = _regras_verticais(texto, [nivel["codigo"] for nivel in niveis])
    for nivel in niveis[1:]:
        regra = regras.get(nivel["codigo"])
        if regra is None:
            acrescimo = round((nivel["valores"][0] / base - 1) * 100, 4)
            regra = {
                "tipo": "percentual",
                "valor": acrescimo,
                "nivel_referencia": niveis[0]["codigo"],
            }
            avisos.append(
                f"O acréscimo do nível {nivel['codigo']} foi inferido matematicamente; confirme na lei."
            )
        nivel["regra_vertical"] = regra
        if regra["tipo"] == "percentual":
            nivel["acrescimo_percentual"] = regra["valor"]
            nivel["acrescimo_valor_fixo"] = None
        else:
            nivel["acrescimo_percentual"] = None
            nivel["acrescimo_valor_fixo"] = regra["valor"]

    avisos.extend(_avisos_vigencia(texto))

    lei_encontrada = re.search(
        r"Anexo\s+(?:[A-Z]|[ÚU]nico)\s+(?:da\s+)?LC\s+\d+/\d{4}",
        texto,
        flags=re.IGNORECASE,
    )
    return {
        "lei": lei_encontrada.group(0) if lei_encontrada else "Não identificada",
        "arquivo_fonte": nome_arquivo,
        "progressao_classes_percentual": progressao if uniforme else None,
        "progressoes_classes_percentuais": progressoes,
        "progressao_classes_uniforme": uniforme,
        "progressao_classes_origem": origem_progressao,
        "classes": list(range(1, quantidade_classes + 1)),
        "niveis": niveis,
        "avisos_importacao": avisos,
    }


def importar_tabela_salarial(nome_arquivo, conteudo):
    texto = extrair_texto_documento(nome_arquivo, conteudo)
    resultado = extrair_tabela_salarial_texto(texto, nome_arquivo=nome_arquivo)
    resultado["texto_extraido_caracteres"] = len(texto)
    return resultado
